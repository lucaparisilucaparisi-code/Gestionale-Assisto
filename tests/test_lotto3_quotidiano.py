"""Lotto 3 "quotidiano": audit delle ore, modifica massiva annullabile, omonimi,
filtri avanzati negli export, parametri di calcolo da config, migrazione completa."""
import io
import json
import os

from openpyxl import load_workbook

PROJECT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def _set_ore(db, uid, anno, mese, ore):
    db.get_or_create_rendicontazione(uid, anno, mese)
    db.update_rendicontazione(uid, anno, mese, ore_lavorate=ore)


def _celle(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if v is not None:
                    yield str(v)


def _scuola(db, commessa):
    db.create_commessa(commessa)
    return db.get_or_create_scuola(commessa, f'IC {commessa} - Primaria')


def _upload(client, contenuto, **form):
    data = {'file': (io.BytesIO(contenuto), 'migrazione.json'), **form}
    return client.post('/api/migrazione/importa', data=data, content_type='multipart/form-data')


# ---------- 12. audit delle ore ----------

def test_modifica_ore_lascia_traccia_nell_audit(client, db_mod):
    db = db_mod
    sid = _scuola(db, 'L3 AUDIT')
    uid = db.get_or_create_utente(sid, 'Audit', 'Ore', 10)
    assert client.post('/api/rendicontazione/2025/11',
                       json={'utente_id': uid, 'ore_lavorate_60': 12}).status_code == 200
    assert client.post('/api/rendicontazione/2025/11/batch',
                       json={'updates': [{'utente_id': uid, 'pasti': 4}]}).status_code == 200

    audit = client.get('/api/audit?entita=rendicontazione&limit=100').get_json()
    azioni = [a for a in audit if a['azione'] == 'modifica_ore']
    assert any(a['entita_id'] == uid and 'ore_lavorate=12' in (a['dettagli'] or '') for a in azioni), \
        'la modifica delle ore non lasciava traccia'
    assert any('1 utenti aggiornati' in (a['dettagli'] or '') for a in azioni)


# ---------- 12. modifica massiva: una transazione, annullabile ----------

def test_bulk_update_tutto_o_niente_e_annullabile(client, db_mod):
    db = db_mod
    sid = _scuola(db, 'L3 BULK')
    a = db.get_or_create_utente(sid, 'Bulk', 'A', 10)
    b = db.get_or_create_utente(sid, 'Bulk', 'B', 12)

    # una riga non valida blocca tutto (prima le altre venivano scritte lo stesso)
    r = client.put('/api/utenti/bulk', json={'updates': [{'id': a, 'monte_ore': 15}, {'id': b, 'monte_ore': 99}]})
    assert r.status_code == 400
    assert db.get_utente_by_id(a)['monte_ore_settimanale'] == 10

    r = client.put('/api/utenti/bulk', json={'updates': [{'id': a, 'monte_ore': 15}, {'id': b, 'lista_attesa': 'Marzo'}]})
    assert r.status_code == 200 and r.get_json()['aggiornati'] == 2
    assert db.get_utente_by_id(a)['monte_ore_settimanale'] == 15
    assert db.get_utente_by_id(b)['lista_attesa'] == 'Marzo'

    assert any(s['tipo'] == 'bulk_update_utenti' for s in client.get('/api/undo').get_json())
    assert client.post('/api/undo').status_code == 200
    assert db.get_utente_by_id(a)['monte_ore_settimanale'] == 10, "l'undo della modifica massiva non ha funzionato"
    assert db.get_utente_by_id(b)['lista_attesa'] is None


def test_bulk_delete_segnala_le_righe_non_riuscite(client, db_mod):
    db = db_mod
    sid = _scuola(db, 'L3 BULKDEL')
    a = db.get_or_create_utente(sid, 'Del', 'A', 10)
    body = client.delete('/api/utenti/bulk', json={'ids': [a, 999999]}).get_json()
    assert body['eliminati'] == 1 and body['success'] is False
    assert body['errori'] and body['errori'][0]['id'] == 999999
    assert db.get_utente_by_id(a) is None


# ---------- 11. omonimi ----------

def test_creazione_omonimo_richiede_conferma(client, db_mod):
    db = db_mod
    _scuola(db, 'L3 OMONIMI')
    payload = {'commessa': 'L3 OMONIMI', 'scuola': 'IC L3 OMONIMI - Primaria',
               'nome': 'Mario', 'cognome': 'Rossi', 'monte_ore': 10}
    assert client.post('/api/utenti', json=payload).status_code == 200

    r = client.post('/api/utenti', json={**payload, 'monte_ore': 20})
    assert r.status_code == 409 and r.get_json()['code'] == 'UTENTE_DUPLICATO'
    primo = r.get_json()['utente_esistente']
    assert primo['monte_ore_settimanale'] == 10, 'il primo utente veniva sovrascritto in silenzio'

    r = client.post('/api/utenti', json={**payload, 'monte_ore': 20, 'forza': True})
    assert r.status_code == 200 and r.get_json()['utente_id'] != primo['id']


# ---------- 4. filtri avanzati negli export ----------

def test_filtri_avanzati_valgono_per_gli_export(client, db_mod):
    db = db_mod
    sid = _scuola(db, 'L3 FILTRI')
    con = db.get_or_create_utente(sid, 'Conore', 'Test', 10)
    senza = db.get_or_create_utente(sid, 'Senzaore', 'Test', 10)
    _set_ore(db, con, 2025, 11, 30)
    _set_ore(db, senza, 2025, 11, 0)

    r = client.get('/api/export/municipale/2025/11?commessa=L3%20FILTRI&ore=zero')
    celle = list(_celle(load_workbook(io.BytesIO(r.data))))
    assert any('Senzaore Test' in c for c in celle) and not any('Conore Test' in c for c in celle)

    r = client.get('/api/export/excel/2025/11?commessa=L3%20FILTRI&search=conore')
    celle = list(_celle(load_workbook(io.BytesIO(r.data))))
    assert any('Conore Test' in c for c in celle) and not any('Senzaore Test' in c for c in celle)


# ---------- 20. parametri di calcolo: unica fonte ----------

def test_parametri_di_calcolo_da_config(client):
    import config
    cfg = client.get('/api/config').get_json()
    assert cfg['tariffa_oraria'] == config.TARIFFA_ORARIA
    assert cfg['iva_percentuale'] == config.IVA_PERCENTUALE
    html = client.get('/rendicontazione').get_data(as_text=True)
    assert 'window.APP_CONFIG' in html and str(config.TARIFFA_ORARIA) in html
    for t in ('rendicontazione', 'index', 'report', 'calendario', 'chiusura_mese'):
        with open(os.path.join(PROJECT_DIR, 'templates', f'{t}.html'), encoding='utf-8') as f:
            src = f.read()
        assert '24.07' not in src and '24,07' not in src, f'tariffa scritta a mano in {t}.html'


# ---------- 17. nessuna finestra di sistema (confirm/prompt) ----------

def test_nessun_confirm_o_prompt_nativo_nelle_pagine():
    import re
    pattern = re.compile(r'(?<![\w.])(confirm|prompt)\(')
    cartelle = [os.path.join(PROJECT_DIR, 'templates'), os.path.join(PROJECT_DIR, 'static', 'js')]
    trovati = []
    for cartella in cartelle:
        for nome in os.listdir(cartella):
            path = os.path.join(cartella, nome)
            if not os.path.isfile(path):
                continue
            with open(path, encoding='utf-8') as f:
                for n, riga in enumerate(f, 1):
                    if pattern.search(riga) and not riga.strip().startswith('//'):
                        trovati.append(f'{nome}:{n}')
    assert not trovati, f'confirm()/prompt() nativi ancora presenti: {trovati}'


# ---------- 13. migrazione completa ----------

def test_migrazione_esporta_tutto_e_sostituisci_ricarica_tutto(client, db_mod):
    db = db_mod
    sid = _scuola(db, 'L3 MIGRA')
    uid = db.get_or_create_utente(sid, 'Migra', 'Test', 10)
    db.add_variazione_monte_ore(uid, 13, '2025-12', 'var')
    did = db.create_dipendente({'nome': 'Op', 'cognome': 'Migra', 'codice_fiscale': 'MIGRA00A01H501X'})
    db.create_assegnazione(uid, did, 10)
    _set_ore(db, uid, 2025, 11, 20)
    client.post('/api/mese-chiuso/2024/1')

    r = client.get('/api/migrazione/esporta')
    assert r.status_code == 200
    export = json.loads(r.data)
    assert export['versione'] == '3.0'
    tab = export['tabelle']
    for t in ('variazioni_monte_ore', 'dipendenti', 'assegnazioni', 'mesi_chiusi', 'impostazioni', 'turni'):
        assert t in tab, f'tabella {t} non esportata'
    assert any(v['utente_id'] == uid and v['monte_ore'] == 13 for v in tab['variazioni_monte_ore'])
    assert any(d['codice_fiscale'] == 'MIGRA00A01H501X' for d in tab['dipendenti'])

    contenuto = json.dumps(export).encode('utf-8')
    anteprima = client.post('/api/migrazione/anteprima',
                            data={'file': (io.BytesIO(contenuto), 'x.json')},
                            content_type='multipart/form-data').get_json()
    assert anteprima['completo'] is True and anteprima['tabelle']['dipendenti'] >= 1

    # sostituisci: senza conferma esplicita non parte
    assert _upload(client, contenuto, mode='replace').status_code == 400

    n_prima = db.count_utenti(attivo=None)
    estraneo = db.get_or_create_utente(sid, 'Estraneo', 'Test', 5)   # non e' nel file: deve sparire
    r = _upload(client, contenuto, mode='replace', confirm='SOSTITUISCI')
    assert r.status_code == 200, r.data
    body = r.get_json()
    assert body['mode'] == 'replace' and body['backup']
    assert db.get_utente_by_id(estraneo) is None
    assert db.count_utenti(attivo=None) == n_prima
    assert db.get_variazioni_monte_ore(uid)[0]['monte_ore'] == 13
    assert len(db.get_assegnazioni_utente(uid)) == 1
    assert db.get_mese_chiuso(2024, 1)
    client.delete('/api/mese-chiuso/2024/1')


def test_migrazione_unisci_importa_anche_le_tabelle_aggiuntive(client, db_mod):
    db = db_mod
    sid = _scuola(db, 'L3 MERGE')
    uid = db.get_or_create_utente(sid, 'Merge', 'Test', 10)
    # file "da un altro PC": stesso utente (id diverso) con una variazione e un operatore assegnato
    file = {
        'versione': '3.0',
        'commesse': [{'id': 77, 'nome': 'L3 MERGE', 'descrizione': None, 'colore': '#000000',
                      'attiva': 1, 'data_creazione': '2025-01-01'}],
        'scuole': [{'id': 88, 'commessa_id': 77, 'commessa_nome': 'L3 MERGE',
                    'nome_completo': 'IC L3 MERGE - Primaria'}],
        'utenti': [{'id': 99, 'scuola_id': 88, 'scuola_nome': 'IC L3 MERGE - Primaria', 'commessa_nome': 'L3 MERGE',
                    'nome': 'Merge', 'cognome': 'Test', 'nome_puntato': 'M. T.', 'monte_ore_settimanale': 11,
                    'lista_attesa': None, 'attivo': 1, 'data_inserimento': '2025-01-01'}],
        'rendicontazione': [], 'calendario': [],
        'tabelle': {
            'dipendenti': [{'id': 5, 'nome': 'Nuovo', 'cognome': 'Operatore', 'codice_fiscale': 'NWOPRT00A01H501Z',
                            'attivo': 1, 'data_inserimento': '2025-01-01'}],
            'variazioni_monte_ore': [{'id': 3, 'utente_id': 99, 'monte_ore': 14, 'mese_inizio': '2026-01',
                                      'mese_fine': None, 'nota': 'dal file', 'data_inserimento': '2025-01-01'}],
            'assegnazioni': [{'id': 9, 'utente_id': 99, 'dipendente_id': 5, 'ore_settimanali': 11,
                              'valido_da': None, 'valido_a': None, 'note': None, 'data_inserimento': '2025-01-01'}],
        },
    }
    contenuto = json.dumps(file).encode('utf-8')
    r = _upload(client, contenuto, mode='merge')
    assert r.status_code == 200, r.data
    stats = r.get_json()['stats']
    assert stats['variazioni_monte_ore']['importati'] == 1
    assert stats['dipendenti']['importati'] == 1 and stats['assegnazioni']['importati'] == 1
    # utente abbinato per nome (monte ore aggiornato); variazione e assegnazione agganciate al suo id reale
    assert db.get_utente_by_id(uid)['monte_ore_settimanale'] == 11
    assert any(v['monte_ore'] == 14 for v in db.get_variazioni_monte_ore(uid))
    assert len(db.get_assegnazioni_utente(uid)) == 1

    # "solo nuovi": niente duplicati e niente modifiche
    db.get_or_create_utente(sid, 'Merge', 'Test', 12)   # riporta il monte ore a 12
    r = _upload(client, contenuto, mode='skip')
    assert r.get_json()['stats']['variazioni_monte_ore']['importati'] == 0
    assert len(db.get_variazioni_monte_ore(uid)) == 1
    assert db.get_utente_by_id(uid)['monte_ore_settimanale'] == 12
