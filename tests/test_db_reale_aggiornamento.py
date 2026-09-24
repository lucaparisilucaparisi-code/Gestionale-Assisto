"""DB reale, parte 2: aggiornamento sicuro dalla versione precedente e rifiniture.

- B1 "Ripristina backup": un backup della versione precedente funziona subito (senza
  riavvio), la pulizia dei vecchi backup non cancella quello scelto, niente ripristino
  senza backup di sicurezza o da un file che non e' un database del gestionale;
- B2 il backup all'avvio si fa PRIMA delle migrazioni (copia "pre-aggiornamento");
- B3 una seconda istanza non parte se Assisto e' gia' aperto sulla porta 5000;
- B4 trasloco JSON "Unisci": un file 2.0 non azzera date/budget/anno della lista,
  niente ore nei mesi chiusi, l'anteprima dice cosa manca;
- B5 README con la guida all'aggiornamento, python-docx >= 1.1.2.

Solo dati inventati: nomi di fantasia e commesse di prova.
"""
import hashlib
import io
import json
import os
import re
import shutil
import socket
import sqlite3
import subprocess
import sys

import pytest

import config

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def _leggi(rel):
    with open(os.path.join(ROOT, rel), encoding='utf-8') as f:
        return f.read()


def _set_ore(db, uid, anno, mese, ore, pasti=None):
    db.get_or_create_rendicontazione(uid, anno, mese)
    db.update_rendicontazione(uid, anno, mese, ore_lavorate=ore, pasti=pasti)


def _q(commessa):
    return commessa.replace(' ', '%20')


def _colonne(percorso, tabella):
    conn = sqlite3.connect(percorso)
    try:
        return [r[1] for r in conn.execute(f'PRAGMA table_info({tabella})')]
    finally:
        conn.close()


def _hash(percorso):
    with open(percorso, 'rb') as f:
        return hashlib.sha256(f.read()).hexdigest()


# ==================== B1: ripristino backup ====================

@pytest.fixture
def db_isolato(db_mod, tmp_path, monkeypatch):
    """Database e cartella dei backup separati da quelli della sessione di test
    (il ripristino sostituisce il database: non deve toccare quello condiviso)."""
    percorso = str(tmp_path / 'gestionale.db')
    cartella = tmp_path / 'backups'
    cartella.mkdir()
    monkeypatch.setattr(db_mod, 'DATABASE_PATH', percorso)
    monkeypatch.setattr(config, 'DATABASE_PATH', percorso)
    monkeypatch.setattr(config, 'BACKUP_FOLDER', str(cartella))
    db_mod.invalida_cache_auth()
    db_mod.init_db()
    from werkzeug.security import generate_password_hash
    db_mod.auth_create_user('prova', generate_password_hash('segreta'))
    yield {'db': percorso, 'backups': str(cartella), 'tmp': tmp_path}
    db_mod.invalida_cache_auth()


def _db_con_dati(db, nome_commessa, nome_utente):
    """Una commessa con un utente, una variazione e le ore di ottobre 2025."""
    db.create_commessa(nome_commessa)
    sid = db.get_or_create_scuola(nome_commessa, 'IC Fantasia - Primaria Girasole')
    uid = db.get_or_create_utente(sid, nome_utente, 'Prova', 10)
    db.add_variazione_monte_ore(uid, 12, '2025-11', 'aumento')
    _set_ore(db, uid, 2025, 10, 30, pasti=5)
    return uid


def _rendi_versione_vecchia(percorso):
    """Toglie dal file le colonne che la versione precedente (1.7.0) non aveva."""
    if sqlite3.sqlite_version_info < (3, 35, 0):
        pytest.skip('SQLite senza ALTER TABLE DROP COLUMN')
    conn = sqlite3.connect(percorso)
    try:
        conn.execute('ALTER TABLE variazioni_monte_ore DROP COLUMN mese_fine')
        conn.execute('ALTER TABLE utenti DROP COLUMN lista_attesa_as')
        conn.execute('ALTER TABLE utenti DROP COLUMN archiviato_dal')
        conn.commit()
    finally:
        conn.close()


def _copia_come_backup(db, cartella, nome):
    """Copia consistente del database attivo nella cartella dei backup."""
    src = sqlite3.connect(db.DATABASE_PATH)
    dst = sqlite3.connect(os.path.join(cartella, nome))
    with dst:
        src.backup(dst)
    src.close()
    dst.close()
    return os.path.join(cartella, nome)


def test_ripristino_di_un_backup_della_versione_precedente_funziona_senza_riavvio(client, db_mod, db_isolato):
    db = db_mod
    uid = _db_con_dati(db, 'RIPR VECCHIO', 'Ripristino')
    vecchio = _copia_come_backup(db, db_isolato['backups'], 'gestionale_backup_20240101_080000.db')
    _rendi_versione_vecchia(vecchio)
    assert 'mese_fine' not in _colonne(vecchio, 'variazioni_monte_ore')
    impronta = _hash(vecchio)

    r = client.post('/api/backup/restore', json={'backup': 'gestionale_backup_20240101_080000.db',
                                                 'confirm': 'CONFERMA'})
    assert r.status_code == 200, r.data
    # lo schema e' gia' aggiornato: niente "no such column" fino al riavvio
    assert 'mese_fine' in _colonne(db_isolato['db'], 'variazioni_monte_ore')
    assert {'lista_attesa_as', 'archiviato_dal'} <= set(_colonne(db_isolato['db'], 'utenti'))
    q = _q('RIPR VECCHIO')
    for url in (f'/api/rendicontazione/2025/10?commessa={q}', f'/api/rendicontazione/2025/11?commessa={q}',
                f'/api/utenti?commessa={q}', f'/api/utenti/{uid}/variazioni-monte-ore',
                f'/api/export/excel/2025/10?commessa={q}', f'/api/export/municipale/2025/10?commessa={q}',
                f'/api/export/annuale/2025-2026?commessa={q}', '/rendicontazione', '/'):
        risposta = client.get(url)
        assert risposta.status_code == 200, (url, risposta.status_code)
    dati = client.get(f'/api/rendicontazione/2025/10?commessa={q}').get_json()['dati']
    assert [(d['utente_id'], d['ore_lavorate_60']) for d in dati] == [(uid, 30)]
    # il file scelto resta com'era (aperto in sola lettura) e non lascia file accanto
    assert _hash(vecchio) == impronta
    assert not [f for f in os.listdir(db_isolato['backups']) if f.endswith(('-wal', '-shm'))]


def test_ripristino_con_cartella_piena_non_cancella_il_backup_scelto(db_mod, db_isolato):
    db = db_mod
    _db_con_dati(db, 'RIPR PIENA', 'Scelto')
    cartella = db_isolato['backups']
    scelto = 'gestionale_backup_20200101_000000.db'          # il piu' vecchio della lista
    _copia_come_backup(db, cartella, scelto)
    # la cartella si riempie con copie di un database diverso (senza l'utente 'Scelto')
    altro = db_isolato['tmp'] / 'altro.db'
    conn = sqlite3.connect(db.DATABASE_PATH)
    dst = sqlite3.connect(str(altro))
    with dst:
        conn.backup(dst)
    conn.close()
    dst.close()
    c = sqlite3.connect(str(altro))
    c.execute("UPDATE utenti SET nome = 'Diverso'")
    c.commit()
    c.close()
    for i in range(config.MAX_BACKUPS + 2):
        shutil.copy(str(altro), os.path.join(cartella, f'gestionale_backup_2021{i + 1:04d}_000000.db'))

    assert db.restore_backup(scelto) is True
    assert os.path.exists(os.path.join(cartella, scelto)), 'la pulizia ha cancellato il backup scelto'
    backups = [f for f in os.listdir(cartella) if f.startswith('gestionale_backup_') and f.endswith('.db')]
    assert len(backups) <= config.MAX_BACKUPS
    with db.get_db_context() as conn:
        nomi = [r[0] for r in conn.execute('SELECT nome FROM utenti')]
    assert nomi == ['Scelto'], 'il database ripristinato non e\' quello scelto'


def test_pulizia_backup_non_cancella_mai_il_file_escluso(db_mod, db_isolato):
    cartella = db_isolato['backups']
    for i in range(config.MAX_BACKUPS + 5):
        open(os.path.join(cartella, f'gestionale_backup_2019{i + 1:04d}_000000.db'), 'wb').close()
    db_mod.cleanup_old_backups(escludi='gestionale_backup_20190001_000000.db')
    rimasti = sorted(os.listdir(cartella))
    assert 'gestionale_backup_20190001_000000.db' in rimasti
    assert len(rimasti) == config.MAX_BACKUPS


def test_niente_ripristino_senza_backup_di_sicurezza(db_mod, db_isolato, monkeypatch):
    db = db_mod
    _db_con_dati(db, 'RIPR SICUREZZA', 'Prima')
    _copia_come_backup(db, db_isolato['backups'], 'gestionale_backup_20240202_000000.db')
    with db.get_db_context() as conn:
        conn.execute("UPDATE utenti SET nome = 'Attuale'")
    monkeypatch.setattr(db, 'create_backup', lambda *a, **k: None)   # disco pieno, permessi...
    assert db.restore_backup('gestionale_backup_20240202_000000.db') is False
    with db.get_db_context() as conn:
        assert [r[0] for r in conn.execute('SELECT nome FROM utenti')] == ['Attuale']


def test_ripristino_rifiuta_file_vuoto_o_estraneo(client, db_mod, db_isolato):
    db = db_mod
    _db_con_dati(db, 'RIPR ESTRANEO', 'Intatto')
    cartella = db_isolato['backups']
    open(os.path.join(cartella, 'gestionale_backup_20240303_000000.db'), 'wb').close()      # vuoto
    estraneo = sqlite3.connect(os.path.join(cartella, 'gestionale_backup_20240304_000000.db'))
    estraneo.execute('CREATE TABLE spesa (voce TEXT)')                                    # altro programma
    estraneo.commit()
    estraneo.close()
    motivi = {'gestionale_backup_20240303_000000.db': 'il file è vuoto',
              'gestionale_backup_20240304_000000.db': 'non è un database del gestionale'}
    for nome, motivo in motivi.items():
        assert db.restore_backup(nome) is False
        r = client.post('/api/backup/restore', json={'backup': nome, 'confirm': 'CONFERMA'})
        assert r.status_code == 400 and 'non sono stati toccati' in r.get_json()['error']
        assert f'Questo file non si può ripristinare: {motivo}' in r.get_json()['error']
        with db.get_db_context() as conn:
            assert [x[0] for x in conn.execute('SELECT nome FROM utenti')] == ['Intatto']


# ==================== B2: backup all'avvio prima delle migrazioni ====================

def _esegui_python(codice, timeout=90):
    # messaggi con le lettere accentate: uscita in UTF-8 qualunque sia la lingua del sistema
    ambiente = dict(os.environ, PYTHONIOENCODING='utf-8')
    return subprocess.run([sys.executable, '-c', codice], capture_output=True, text=True,
                          encoding='utf-8', errors='replace', timeout=timeout, cwd=ROOT, env=ambiente)


def test_backup_all_avvio_prima_delle_migrazioni(db_mod, db_isolato):
    db = db_mod
    _db_con_dati(db, 'AVVIO VECCHIO', 'Avvio')
    _rendi_versione_vecchia(db_isolato['db'])
    cartella = db_isolato['backups']
    codice = f'''
import sys
sys.path.insert(0, {ROOT!r})
import config
config.DATABASE_PATH = {db_isolato['db']!r}
config.BACKUP_FOLDER = {cartella!r}
config.BACKUP_ON_STARTUP = True
import database
database.DATABASE_PATH = config.DATABASE_PATH
import app
'''
    esito = _esegui_python(codice)
    assert esito.returncode == 0, esito.stderr[-2000:]
    backups = [f for f in os.listdir(cartella) if f.endswith('.db')]
    assert len(backups) == 1
    # la copia e' com'era PRIMA dell'aggiornamento, il database attivo e' aggiornato
    assert 'mese_fine' not in _colonne(os.path.join(cartella, backups[0]), 'variazioni_monte_ore')
    assert 'mese_fine' in _colonne(db_isolato['db'], 'variazioni_monte_ore')


# ==================== B3: doppia istanza ====================

def test_assisto_gia_aperto_riconosce_la_porta_occupata(app_module):
    s = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
    s.bind(('127.0.0.1', 0))
    s.listen(1)
    porta = s.getsockname()[1]
    try:
        assert app_module.assisto_gia_aperto(porta) is True
    finally:
        s.close()
    assert app_module.assisto_gia_aperto(porta) is False
    # il controllo sta nel blocco __main__, prima di app.run
    main = _leggi('app.py').split("if __name__ == '__main__':", 1)[1]
    assert main.index('assisto_gia_aperto()') < main.index('app.run(')
    assert "Assisto è già aperto" in app_module.MESSAGGIO_GIA_APERTO


def test_seconda_istanza_esce_con_messaggio_chiaro(tmp_path):
    """Avvio vero di app.py come programma con la porta 5000 gia' occupata."""
    occupante = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
    try:
        occupante.bind(('127.0.0.1', 5000))
        occupante.listen(1)
    except OSError:
        occupante.close()
        occupante = None          # la porta e' gia' occupata da un altro programma: va bene uguale
    codice = f'''
import runpy, sys
sys.path.insert(0, {ROOT!r})
import config
config.DATABASE_PATH = {str(tmp_path / 'g.db')!r}
config.BACKUP_FOLDER = {str(tmp_path / 'backups')!r}
config.BACKUP_ON_STARTUP = False
import database
database.DATABASE_PATH = config.DATABASE_PATH
runpy.run_path({os.path.join(ROOT, 'app.py')!r}, run_name='__main__')
'''
    try:
        esito = _esegui_python(codice)
    finally:
        if occupante:
            occupante.close()
    assert esito.returncode == 1
    assert "Assisto è già aperto" in esito.stdout + esito.stderr


# ==================== B4: trasloco JSON ====================

def _file_json(commessa, utenti, rendicontazione=(), versione='2.0', tabelle=None):
    dati = {'versione': versione,
            'commesse': [{'id': 501, 'nome': commessa, 'descrizione': None, 'colore': '#3B82F6', 'attiva': 1}],
            'scuole': [{'id': 502, 'commessa_id': 501, 'commessa_nome': commessa,
                        'nome_completo': 'IC Fantasia - Primaria Trasloco'}],
            'utenti': list(utenti), 'rendicontazione': list(rendicontazione), 'calendario': []}
    if tabelle is not None:
        dati['tabelle'] = tabelle
    return json.dumps(dati).encode('utf-8')


def _utente_v170(uid_file, nome, monte_ore=10, lista=None, attivo=1):
    """Utente come lo esporta la versione 1.7.0 (formato 2.0): senza date di
    servizio, budget, anno della lista e mese dell'archiviazione."""
    return {'id': uid_file, 'scuola_id': 502, 'scuola_nome': 'IC Fantasia - Primaria Trasloco',
            'commessa_nome': None, 'nome': nome, 'cognome': 'Trasloco', 'nome_puntato': f'{nome[0]}. T.',
            'monte_ore_settimanale': monte_ore, 'lista_attesa': lista, 'attivo': attivo,
            'data_inserimento': '2025-09-01T10:00:00'}


def _registro_migrazioni(db):
    return [v for v in db.get_audit_log(limit=50, entita='sistema') if v['azione'] == 'migrazione']


def _importa(client, contenuto, mode='merge'):
    return client.post('/api/migrazione/importa', data={
        'mode': mode, 'file': (io.BytesIO(contenuto), 'trasloco.json')}, content_type='multipart/form-data')


def test_unisci_file_vecchio_non_azzera_date_budget_e_anno_lista(client, db_mod):
    db = db_mod
    commessa = 'TRASLOCO CAMPI'
    db.create_commessa(commessa)
    sid = db.get_or_create_scuola(commessa, 'IC Fantasia - Primaria Trasloco')
    uid = db.get_or_create_utente(sid, 'Datato', 'Trasloco', 10)
    arch = db.get_or_create_utente(sid, 'Archiviato', 'Trasloco', 8)
    db.update_utente_periodo(uid, '2025-10', '2026-05')
    db.update_utente_lista_attesa(uid, 'Marzo', '2025-2026')
    with db.get_db_context() as conn:
        conn.execute('UPDATE utenti SET budget_ore_mensile = 40, budget_ore_annuale = 400 WHERE id = ?', (uid,))
        conn.execute("UPDATE utenti SET attivo = 0, archiviato_dal = '2026-02' WHERE id = ?", (arch,))
    utenti = [_utente_v170(801, 'Datato', monte_ore=11, lista='Marzo'),
              _utente_v170(802, 'Archiviato', monte_ore=8, attivo=0)]
    for u in utenti:
        u['commessa_nome'] = commessa
    r = _importa(client, _file_json(commessa, utenti))
    assert r.status_code == 200, r.data
    assert r.get_json()['stats']['utenti']['aggiornati'] == 2

    u = db.get_utente_by_id(uid)
    assert u['monte_ore_settimanale'] == 11                      # campo presente nel file: aggiornato
    assert (u['data_inizio'], u['data_fine']) == ('2025-10', '2026-05')
    assert (u['budget_ore_mensile'], u['budget_ore_annuale']) == (40, 400)
    assert (u['lista_attesa'], u['lista_attesa_as']) == ('Marzo', '2025-2026')
    a = db.get_utente_by_id(arch)
    assert (a['attivo'], a['archiviato_dal']) == (0, '2026-02')


def test_unisci_non_scrive_ore_nei_mesi_chiusi_e_lo_dice(client, db_mod):
    db = db_mod
    commessa = 'TRASLOCO CHIUSO'
    db.create_commessa(commessa)
    sid = db.get_or_create_scuola(commessa, 'IC Fantasia - Primaria Trasloco')
    uid = db.get_or_create_utente(sid, 'Chiuso', 'Trasloco', 10)
    nuovo = db.get_or_create_utente(sid, 'Senzariga', 'Trasloco', 10)
    _set_ore(db, uid, 2024, 6, 10.5, pasti=4)
    _set_ore(db, uid, 2024, 5, 20)
    db.set_mese_chiuso(2024, 6, True)
    try:
        def riga(i, u_file, nome, mese, ore, pasti=0):
            return {'id': i, 'utente_id': u_file, 'utente_nome': nome, 'utente_cognome': 'Trasloco',
                    'anno': 2024, 'mese': mese, 'ore_lavorate_60': ore, 'pasti': pasti,
                    'giorni_lavorativi': 20, 'note': None}
        utenti = [_utente_v170(811, 'Chiuso'), _utente_v170(812, 'Senzariga')]
        righe = [riga(1, 811, 'Chiuso', 6, 9.5), riga(2, 812, 'Senzariga', 6, 7),   # giugno chiuso
                 riga(3, 811, 'Chiuso', 5, 22)]                                     # maggio aperto
        r = _importa(client, _file_json(commessa, utenti, righe))
        assert r.status_code == 200, r.data
        esito = r.get_json()
        assert esito['stats']['rendicontazione']['saltate_mese_chiuso'] == 2
        assert esito['mesi_chiusi_saltati'] == [{'anno': 2024, 'mese': 6}]
        assert esito['avvisi'] == ['2 righe di ore NON sono state scritte perché il mese è chiuso (Giugno 2024). '
                                   "Per cambiarle riapri il mese da Chiusura Mese e ripeti l'importazione."]
        assert any(v['dettagli'].endswith('; saltate 2 righe di mesi chiusi') for v in _registro_migrazioni(db))
        with db.get_db_context() as conn:
            ore = {(row[0], row[1]): row[2] for row in conn.execute(
                'SELECT utente_id, mese, ore_lavorate_60 FROM rendicontazione WHERE utente_id IN (?, ?)',
                (uid, nuovo))}
        assert ore == {(uid, 6): 10.5, (uid, 5): 22}                 # giugno intatto, niente riga nuova

        # un file con la stessa riga del mese chiuso non produce avvisi
        with db.get_db_context() as conn:
            giorni = conn.execute('SELECT giorni_lavorativi FROM rendicontazione WHERE utente_id = ? '
                                  'AND anno = 2024 AND mese = 6', (uid,)).fetchone()[0]
        righe_uguali = [dict(riga(1, 811, 'Chiuso', 6, 10.5, pasti=4), giorni_lavorativi=giorni)]
        esito = _importa(client, _file_json(commessa, utenti[:1], righe_uguali)).get_json()
        assert esito['stats']['rendicontazione']['saltate_mese_chiuso'] == 0 and not esito['avvisi']

        # una sola riga saltata: al singolare (prima "1 righe di ore NON sono state scritte")
        esito = _importa(client, _file_json(commessa, utenti[:1], [riga(1, 811, 'Chiuso', 6, 9.5)])).get_json()
        assert esito['avvisi'] == ['1 riga di ore NON è stata scritta perché il mese è chiuso (Giugno 2024). '
                                   "Per cambiarla riapri il mese da Chiusura Mese e ripeti l'importazione."]
        assert any(v['dettagli'].endswith('; saltata 1 riga di mesi chiusi') for v in _registro_migrazioni(db))
    finally:
        db.set_mese_chiuso(2024, 6, False)


def test_anteprima_file_vecchio_elenca_cosa_manca(client):
    commessa = 'TRASLOCO ANTEPRIMA'
    vecchio = _file_json(commessa, [_utente_v170(821, 'Anteprima')])
    r = client.post('/api/migrazione/anteprima', data={'file': (io.BytesIO(vecchio), 'v170.json')},
                    content_type='multipart/form-data')
    j = r.get_json()
    testo = ' '.join(j['mancanti'])
    for voce in ('determine', 'recuperi', 'correzioni', 'date di inizio e fine', 'storico'):
        assert voce in testo, voce
    assert 'gestionale.db' in j['consiglio'] and 'Così passano tutti i dati' in j['consiglio']
    assert 'storico delle modifiche (registro attività)' in j['mancanti']
    completo = _file_json(commessa, [dict(_utente_v170(821, 'Anteprima'), data_inizio=None)],
                          versione='3.0', tabelle={'mesi_chiusi': []})
    j = client.post('/api/migrazione/anteprima', data={'file': (io.BytesIO(completo), 'v3.json')},
                    content_type='multipart/form-data').get_json()
    assert j['mancanti'] == [] and j['consiglio'] is None
    # la pagina mostra l'elenco e chiede conferma prima di importare un file incompleto
    pagina = _leggi('templates/import.html')
    assert 'data.mancanti' in pagina and "'Importa comunque'" in pagina


# ==================== B5: guida all'aggiornamento ====================

def test_readme_spiega_come_aggiornare():
    readme = _leggi('README.md')
    assert '## Aggiornare a una nuova versione' in readme
    sezione = readme.split('## Aggiornare a una nuova versione', 1)[1].split('\n## ', 1)[0]
    for passo in ('Chiudi il programma', 'gestionale.db-wal', 'cartella nuova', 'Prima del primo avvio',
                  'backups', 'uploads', 'avvia.bat', 'Tieni la cartella vecchia', 'Trasloco'):
        assert passo in sezione, passo
    assert 'gestionale-oepac' not in readme                          # nome della cartella dello zip
    assert 'Gestionale-Assisto-' in readme
    assert re.search(r'python-docx>=1\.1\.2', _leggi('requirements.txt'))


# ==================== C: Rendicontazione ====================

def _rend_pagina():
    return _leggi('templates/rendicontazione.html')


def test_c1_uscita_dalla_pagina_invia_le_modifiche_in_sospeso():
    """Solo-browser (verifica Playwright nel rapporto): qui si controlla che la pagina
    abbia l'invio keepalive su pagehide/visibilitychange verso il batch del mese mostrato."""
    pagina = _rend_pagina()
    assert "addEventListener('pagehide', inviaModificheInUscita)" in pagina
    assert "visibilitychange" in pagina and "keepalive: true" in pagina
    corpo = pagina.split('function inviaModificheInUscita()', 1)[1].split('\n}\n', 1)[0]
    assert "periodoCaricato" in corpo and '/batch' in corpo
    assert "table-input.changed" in corpo                     # niente invio senza caselle modificate
    # lo stesso pacchetto non parte due volte (scheda nascosta e poi chiusa)...
    assert "if (ultimoInvioUscita === firma) return;" in corpo
    # ...ma ogni nuova modifica azzera la firma: 7:15 (scheda nascosta), 10:00 (salvato
    # in automatico), di nuovo 7:15 e uscita entro 2 s: prima restava 10:00 nel database
    segna = pagina.split('function markChanged(input) {', 1)[1].split('\n}\n', 1)[0]
    assert "ultimoInvioUscita = null;" in segna
    # "Aggiorna" prima salva
    assert "getElementById('btn-load').addEventListener('click', cambiaPeriodo)" in pagina


def _commessa_con_utente(db, commessa, scuola='IC Fantasia - Primaria Limiti', nome='Limite'):
    db.create_commessa(commessa)
    sid = db.get_or_create_scuola(commessa, scuola)
    return sid, db.get_or_create_utente(sid, nome, 'Prova', 10)


def test_c2_pasti_fino_a_62_e_errore_con_utente_e_campo(client, db_mod):
    db = db_mod
    assert config.MAX_PASTI_MENSILI == 62
    _, uid = _commessa_con_utente(db, 'LIMITI PASTI')
    _, altro = _commessa_con_utente(db, 'LIMITI PASTI', nome='Buono')
    base = '/api/rendicontazione/2033/10/batch'
    assert client.post(base, json={'updates': [{'utente_id': uid, 'pasti': 44}]}).status_code == 200
    r = client.post(base, json={'updates': [{'utente_id': altro, 'ore_lavorate_60': 10},
                                            {'utente_id': uid, 'pasti': 63}]})
    assert r.status_code == 400
    j = r.get_json()
    assert 'Limite Prova' in j['error'] and 'Pasti' in j['error']
    assert (j['utente_id'], j['campo']) == (uid, 'pasti')
    r = client.post(base, json={'updates': [{'utente_id': uid, 'ore_lavorate_60': 1230}]})
    assert r.get_json()['campo'] == 'ore_lavorate_60'


def test_c2_pagina_controlla_le_caselle_prima_di_inviare():
    pagina = _rend_pagina()
    assert 'max_pasti_mensili' in pagina and 'max_ore_mensili' in pagina
    assert 'massimo ${MAX_PASTI_MESE} pasti' in pagina
    assert 'forse volevi scrivere' in pagina
    raccolta = pagina.split('function collectPendingChangesFromInputs()', 1)[1].split('\n}\n', 1)[0]
    assert 'controllaCasella(input)' in raccolta and 'segnaCasellaNonValida' in raccolta
    salva = pagina.split('async function autoSave()', 1)[1].split('\n}\n', 1)[0]
    assert 'error.message' in salva and 'd.utente_id && d.campo' in salva


def test_c2_validazione_segnala_pasti_oltre_i_giorni_di_scuola(client, db_mod):
    db = db_mod
    db.set_calendario('2033-2034', 3, 2034, 20, 20)
    _, uid = _commessa_con_utente(db, 'LIMITI GIORNI', nome='Pasti')
    _set_ore(db, uid, 2034, 3, 44, pasti=44)
    j = client.get(f'/api/stats/validazione?anno=2034&mese=3&commessa={_q("LIMITI GIORNI")}').get_json()
    anomalia = [a for a in j['anomalie'] if a['categoria'] == 'pasti_oltre_giorni']
    assert anomalia and anomalia[0]['tipo'] == 'warning'                 # non bloccante
    assert anomalia[0]['titolo'] == 'Pasti più dei giorni di scuola'
    assert anomalia[0]['messaggio'].startswith('1 utente ha più pasti dei giorni di scuola')
    assert anomalia[0]['dettagli'][0]['id'] == uid and anomalia[0]['dettagli'][0]['giorni'] == 20
    db.update_rendicontazione(uid, 2034, 3, pasti=18)
    j = client.get(f'/api/stats/validazione?anno=2034&mese=3&commessa={_q("LIMITI GIORNI")}').get_json()
    assert not [a for a in j['anomalie'] if a['categoria'] == 'pasti_oltre_giorni']


def test_c2_import_excel_con_gli_stessi_limiti(client, db_mod):
    from openpyxl import Workbook
    db = db_mod
    _, buono = _commessa_con_utente(db, 'LIMITI EXCEL', nome='Excelbuono')
    _, pasti = _commessa_con_utente(db, 'LIMITI EXCEL', nome='Excelpasti')
    _, ore = _commessa_con_utente(db, 'LIMITI EXCEL', nome='Excelore')
    wb = Workbook()
    ws = wb.active
    ws.title = 'Estratto Novembre'
    ws.append(['ESTRATTO NOVEMBRE 2033'])
    ws.append(['Commessa', 'Scuola', 'Utenti', "Totale ore lavorate in 100'", 'Pasti'])
    scuola = 'IC Fantasia - Primaria Limiti'
    ws.append(['LIMITI EXCEL', scuola, 'Excelbuono Prova', 30, 10])
    ws.append(['LIMITI EXCEL', scuola, 'Excelpasti Prova', 30, 70])
    ws.append(['LIMITI EXCEL', scuola, 'Excelore Prova', 1230, 5])
    buf = io.BytesIO()
    wb.save(buf)
    contenuto = buf.getvalue()
    anteprima = client.post('/api/import-rendicontazione/preview',
                            data={'file': (io.BytesIO(contenuto), 'limiti.xlsx')},
                            content_type='multipart/form-data').get_json()
    assert anteprima['totale_fuori_limite'] == 2
    r = client.post('/api/import-rendicontazione',
                    data={'file': (io.BytesIO(contenuto), 'limiti.xlsx'), 'modalita': 'overwrite'},
                    content_type='multipart/form-data')
    j = r.get_json()
    assert j['totale_fuori_limite'] == 2
    motivi = ' '.join(v['motivo'] for d in j['dettaglio'] for v in d.get('fuori_limite', []))
    assert 'pasti 70' in motivi and 'ore 1230' in motivi
    assert 'pasti 70: il massimo è 62' in motivi
    righe = {d['utente_id']: (d['ore_lavorate_60'], d['pasti'])
             for d in db.get_rendicontazione_completa(2033, 11, 'LIMITI EXCEL')}
    assert righe[buono] == (30, 10)
    assert not righe[pasti][0] and not righe[ore][0]                # fuori limite: non scritte


def test_c3_totale_plesso_per_id_con_scuole_omonime(client, db_mod):
    db = db_mod
    nome = 'IC Omonimo - Primaria Stesso Nome'
    for commessa, ore in (('OMONIMI A', [30, 20]), ('OMONIMI B', [7])):
        db.create_commessa(commessa)
        sid = db.get_or_create_scuola(commessa, nome)
        for i, h in enumerate(ore):
            uid = db.get_or_create_utente(sid, f'Omonimo{commessa[-1]}{i}', 'Prova', 10)
            _set_ore(db, uid, 2033, 12, h)
    j = client.get('/api/rendicontazione/2033/12').get_json()
    totali = [t for t in j['totali_scuola'] if t['scuola'] == nome]
    assert len(totali) == 2 and all('scuola_id' in t for t in totali)
    per_commessa = {t['commessa']: (t['num_utenti'], t['ore_lavorate_60']) for t in totali}
    assert per_commessa == {'OMONIMI A': (2, 50), 'OMONIMI B': (1, 7)}
    pagina = _rend_pagina()
    assert 't.scuola_id === group.scuola_id' in pagina
    totali_js = pagina.split('function calculateTotaliScuola(dati)', 1)[1].split('\n}\n', 1)[0]
    assert 'grouped[d.scuola_id]' in totali_js and 'grouped[d.scuola]' not in totali_js
    assert 'String(d.scuola_id) === scuolaFilter' in pagina              # filtro Scuola per plesso
    # Excel senza filtro commessa: due blocchi, con la commessa nell'intestazione
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(client.get('/api/export/excel/2033/12').data))
    intestazioni = [c for r in wb['Dettaglio per Scuola'].iter_rows(values_only=True) for c in r
                    if isinstance(c, str) and c.startswith('⊟') and 'Stesso Nome' in c]
    assert sorted(intestazioni) == [f'⊟ OMONIMI A - {nome}', f'⊟ OMONIMI B - {nome}']
    # esporta filtrati di un plesso preciso (scuola_id)
    sid_b = [t['scuola_id'] for t in totali if t['commessa'] == 'OMONIMI B'][0]
    wb = load_workbook(io.BytesIO(client.get(f'/api/export/excel/2033/12?scuola_id={sid_b}').data))
    celle = [c for r in wb['Dettaglio per Scuola'].iter_rows(values_only=True) for c in r if isinstance(c, str)]
    assert any('OmonimoB0' in c for c in celle) and not any('OmonimoA0' in c for c in celle)


def test_c4_copia_mese_prec_avvisa_se_i_giorni_sono_molto_diversi(client, db_mod):
    db = db_mod
    db.set_calendario('2033-2034', 6, 2034, 21, 5)          # giugno: infanzia 21, altri 5
    db.set_calendario('2034-2035', 9, 2034, 12, 12)
    db.set_calendario('2034-2035', 10, 2034, 22, 22)
    db.set_calendario('2034-2035', 11, 2034, 20, 20)
    commessa = 'COPIA GIORNI'
    db.create_commessa(commessa)
    for scuola in ('IC Fantasia - Primaria Copia', "IC Fantasia - Scuola dell'Infanzia Copia"):
        sid = db.get_or_create_scuola(commessa, scuola)
        db.get_or_create_utente(sid, 'Copia' + scuola[-12:-6].strip(), 'Prova', 10)
    q = _q(commessa)
    j = client.get(f'/api/rendicontazione/2034/9/copia-precedente?commessa={q}').get_json()
    assert j['giorni_diversi'] is True and j['mese_origine'] == 'Giugno 2034'
    giorni = {g['tipo']: (g['origine'], g['destinazione']) for g in j['giorni']}
    assert giorni == {'altri': (5, 12), 'infanzia': (21, 12)}
    assert client.get(f'/api/rendicontazione/2034/10/copia-precedente?commessa={q}').get_json()['giorni_diversi']
    assert client.get(f'/api/rendicontazione/2034/11/copia-precedente?commessa={q}').get_json()['giorni_diversi'] is False
    pagina = _rend_pagina()
    assert "confirmText: 'Copia comunque'" in pagina and "extraText: 'Compila con media'" in pagina


def test_c5_riepilogo_usa_il_valore_numerico_salvato():
    pagina = _rend_pagina()
    assert 'data-ore="${row.ore_lavorate_60 || 0}"' in pagina
    totali = pagina.split('function recalculateTotals()', 1)[1].split('\n}\n', 1)[0]
    assert 'dataset.ore' in totali and 'parseTimeInput' not in totali
    assert 'calcolaFatturazione(' in totali
    assert 'input.dataset.ore = ore60' in pagina


# Dopo una modifica la pagina ricalcola riga e Riepilogo da sola: devono coincidere al
# centesimo con la pagina ricaricata (server). Prima: Math.round(v * 100) / 100 (Riepilogo
# di novembre 632.507,51 contro 632.507,50), importi della riga dalle ore gia' arrotondate
# (10:40 -> 256,83 contro 256,75), somma semplice anche col Python 3.12 del server
# (somma compensata: un centesimo in circa 3 Riepiloghi su 100 dei casi con ore ,50).

def _corpo_js(pagina, nome):
    return pagina.split(f'function {nome}(', 1)[1].split('\n}\n', 1)[0]


def _funzioni_calcolo_pagina():
    """Le funzioni di calcolo della Rendicontazione, estratte dal template."""
    pagina = _rend_pagina()
    pezzi = []
    for nome in ('arrotonda2', 'sommaComePython', 'inOrdineServer', 'calcolaFatturazione'):
        inizio = pagina.index(f'function {nome}(')
        pezzi.append(pagina[inizio:pagina.index('\n}\n', inizio) + 2])
    return '\n'.join(pezzi)


def _node_pagina(cfg, programma, dati):
    """Esegue `programma` (corpo di funzione che ritorna un valore JSON) in node, con le
    funzioni di calcolo della pagina, CFG come window.APP_CONFIG e DATI."""
    node = shutil.which('node')
    if not node:
        pytest.skip('node non disponibile: calcolo della pagina non verificabile qui')
    codice = (f'const CFG = {json.dumps(cfg)};\n{_funzioni_calcolo_pagina()}\n'
              "const DATI = JSON.parse(require('fs').readFileSync(0, 'utf8'));\n"
              f'process.stdout.write(JSON.stringify((() => {{ {programma} }})()));')
    r = subprocess.run([node, '-e', codice], input=json.dumps(dati), capture_output=True,
                       text=True, timeout=120)
    assert r.returncode == 0, r.stderr
    return json.loads(r.stdout)


def _somma_semplice(valori):
    somma = 0.0
    for x in valori:
        somma += x
    return somma


def _somma_neumaier(valori):
    """sum() di CPython sui float dalla 3.12."""
    somma = compenso = 0.0
    for x in valori:
        t = somma + x
        compenso += (somma - t) + x if abs(somma) >= abs(x) else (x - t) + somma
        somma = t
    return somma + compenso if compenso else somma


def _diversi(a, b):
    return [(i, x, y) for i, (x, y) in enumerate(zip(a, b)) if x != y][:5]


def test_c5_pagina_ricalcola_come_il_server(client):
    pagina = _rend_pagina()
    fatt = _corpo_js(pagina, 'calcolaFatturazione')
    assert 'Math.round' not in fatt and 'arrotonda2(' in fatt
    arr = _corpo_js(pagina, 'arrotonda2')
    assert 'Math.round' not in arr and 'toFixed(2)' in arr
    # importi della riga dalle ore NON arrotondate, come get_rendicontazione_completa
    ore = _corpo_js(pagina, 'handleOreChange')
    assert 'calcolaFatturazione(ore60)' in ore and 'calcolaFatturazione(ore100)' not in ore
    assert 'const ore100 = arrotonda2(ore60)' in ore and 'Math.round' not in ore
    # la riga in memoria segue la casella (filtro Scuola: prima tornavano i valori del caricamento)
    assert 'Object.assign(utente, {' in ore
    assert 'if (utente) utente.pasti = pasti;' in _corpo_js(pagina, 'handlePastiChange')
    # Riepilogo: round(x, 2) riga per riga, sum() di Python nell'ordine del server
    tot = _corpo_js(pagina, 'recalculateTotals')
    assert 'Math.round' not in tot
    assert 'sommaComePython(inOrdine.map(r => arrotonda2(r.ore60)))' in tot
    assert 'inOrdineServer(righe)' in tot and 'row.dataset.ordine' in tot
    assert 'sommaComePython(inOrdineServer(dati)' in _corpo_js(pagina, 'calculateTotaleGenerale')
    assert 'inOrdineServer(dati).forEach' in _corpo_js(pagina, 'calculateTotaliScuola')
    carica = _corpo_js(pagina, 'loadData')
    assert carica.index('d.ordine_server = i') < carica.index('data.dati.sort(')
    assert 'data-ordine="${row.ordine_server || 0}"' in pagina
    # un valore che parte prima dell'evento change (stesso testo riscritto) ricalcola la riga
    raccolta = _corpo_js(pagina, 'collectPendingChangesFromInputs')
    assert 'esito.valore !== parseFloat(input.dataset.ore)' in raccolta
    assert 'handleOreChange(input)' in raccolta
    # il server dice come somma sum() sul suo Python
    assert config.SOMMA_COMPENSATA == (sys.version_info >= (3, 12))
    assert client.get('/api/config').get_json()['somma_compensata'] is config.SOMMA_COMPENSATA
    assert 'somma_compensata' in client.get('/rendicontazione').get_data(as_text=True)


def test_c5_arrotondamento_e_somme_della_pagina_uguali_a_python():
    import random
    rnd = random.Random(1311)
    tariffa, iva = config.TARIFFA_ORARIA, config.IVA_PERCENTUALE
    # 30119.405: Math.round(v * 100) / 100 da' 30119,41, Python 30119,40; x,125 e simili:
    # pareggi esatti, Python va al centesimo pari (toFixed andrebbe in su)
    valori = [30119.405, 6920.125, 12.625, 0.125, 0.375, 0.625, 0.875, 2.675, 1.005, 0.0]
    valori += [rnd.randint(0, 3_000_000) / 100 * tariffa for _ in range(4000)]
    valori += [rnd.randint(0, 3_000_000) / 100 * iva for _ in range(4000)]
    valori += [rnd.randint(0, 8_000_000) / 8 for _ in range(1000)]
    ore = [rnd.randint(0, 199) + rnd.randint(0, 59) / 60 for _ in range(3000)]   # H:MM scritte
    ore += [rnd.randint(0, 20000) / 100 for _ in range(3000)]                     # decimali importati
    mesi = [[rnd.choice(ore) if rnd.random() < 0.9 else 0.0 for _ in range(rnd.randint(20, 700))]
            for _ in range(400)]
    programma = '''
        return {
            arrotondati: DATI.valori.map(arrotonda2),
            righe: DATI.ore.map(o => { const f = calcolaFatturazione(o); return [arrotonda2(o), f.imponibile, f.iva, f.totale]; }),
            totali: DATI.mesi.map(m => calcolaFatturazione(sommaComePython(m.map(arrotonda2))).totale),
        };'''
    dati = {'valori': valori, 'ore': ore, 'mesi': mesi}
    for compensata, somma in ((False, _somma_semplice), (True, _somma_neumaier)):
        cfg = {'tariffa_oraria': tariffa, 'iva_percentuale': iva, 'somma_compensata': compensata}
        js = _node_pagina(cfg, programma, dati)
        assert not _diversi(js['arrotondati'], [round(v, 2) for v in valori])
        assert not _diversi(js['righe'], [[round(o, 2), *config.calcola_fatturazione(o)] for o in ore])
        attesi = [config.calcola_fatturazione(somma([round(x, 2) for x in m]))[2] for m in mesi]
        assert not _diversi(js['totali'], attesi)
    # il riferimento in Python della somma di questo interprete e' proprio sum()
    somma_qui = _somma_neumaier if config.SOMMA_COMPENSATA else _somma_semplice
    assert all(somma_qui([round(x, 2) for x in m]) == sum(round(x, 2) for x in m) for m in mesi)


def test_c5_riepilogo_e_righe_della_pagina_uguali_al_server(client, db_mod):
    """Mese inventato: il Riepilogo e gli importi di riga che la pagina ricalcola (con le
    sue funzioni, righe nell'ordine della tabella) coincidono con /api/rendicontazione,
    anche dopo modifiche salvate come fa la pagina (H:MM -> ore decimali)."""
    import random
    db = db_mod
    rnd = random.Random(5)
    anno, mese, commessa = 2039, 11, 'CENTESIMI C5'
    db.create_commessa(commessa)
    uids = []
    for s in range(4):
        sid = db.get_or_create_scuola(commessa, f'IC Fantasia - Primaria Centesimi {s}')
        for i in range(45):
            uid = db.get_or_create_utente(sid, f'Cent{s}x{i}', 'Prova', rnd.choice([6, 10, 12, 18]))
            ore = rnd.choice([0, rnd.randint(0, 12000) / 100, rnd.randint(0, 99) + rnd.randint(0, 59) / 60])
            _set_ore(db, uid, anno, mese, ore)
            uids.append(uid)
    cfg = client.get('/api/config').get_json()
    url = f'/api/rendicontazione/{anno}/{mese}?commessa={_q(commessa)}'
    programma = '''
        const righe = DATI.dati.map((d, i) => ({ ordine_server: i, ore60: d.ore_lavorate_60 || 0 })).reverse();
        const tot = calcolaFatturazione(sommaComePython(inOrdineServer(righe).map(r => arrotonda2(r.ore60)))).totale;
        return { tot, righe: DATI.dati.map(d => { const o = d.ore_lavorate_60 || 0; const f = calcolaFatturazione(o);
                                                  return [arrotonda2(o), f.imponibile, f.iva, f.totale]; }) };'''
    confronti = 0
    for giro in range(25):
        j = client.get(url).get_json()
        js = _node_pagina(cfg, programma, {'dati': j['dati']})
        assert js['tot'] == j['totale_generale']['totale_100'], giro
        server = [[d['ore_lavorate_100'], d['imponibile_100'], d['iva_100'], d['totale_100']] for d in j['dati']]
        assert not _diversi(js['righe'], server), giro
        confronti += 1
        # modifica come la pagina: H:MM scritto -> ore + minuti / 60
        uid = rnd.choice(uids)
        ore = rnd.randint(1, 60) + rnd.choice([10, 20, 40, 50, 5, 25, 35]) / 60
        r = client.post(f'/api/rendicontazione/{anno}/{mese}/batch',
                        json={'updates': [{'utente_id': uid, 'ore_lavorate_60': ore}]})
        assert r.status_code == 200
    assert confronti == 25


def test_c6_risposte_fuori_ordine_scartate():
    pagina = _rend_pagina()
    carica = pagina.split('async function loadData()', 1)[1].split('\n}\n', 1)[0]
    assert 'const mio = ++seqCaricamento' in carica
    assert carica.index('if (mio !== seqCaricamento) return;') < carica.index('allData = data.dati')


# ==================== D: Statistiche e Dashboard ====================

def test_d1_heatmap_usa_i_giorni_del_tipo_di_scuola(client, db_mod):
    db = db_mod
    db.set_calendario('2035-2036', 6, 2036, 21, 5)          # giugno: infanzia 21, altri ordini 5
    commessa = 'HEATMAP GIUGNO'
    db.create_commessa(commessa)
    primaria = db.get_or_create_scuola(commessa, 'IC Fantasia - Primaria Heatmap')
    infanzia = db.get_or_create_scuola(commessa, "IC Fantasia - Scuola dell'Infanzia Heatmap")
    u_pri = db.get_or_create_utente(primaria, 'Heatprimaria', 'Prova', 10)
    u_inf = db.get_or_create_utente(infanzia, 'Heatinfanzia', 'Prova', 10)
    for uid, giorni in ((u_pri, 5), (u_inf, 21)):
        _, previste = db.calcola_media_prevista(10, giorni)
        _set_ore(db, uid, 2036, 6, round(previste, 2))
    j = client.get(f'/api/stats/heatmap/2035-2036?commessa={_q(commessa)}&limit=500').get_json()
    giugno = {r['id']: [m for m in r['mesi'] if m['mese'] == 6][0] for r in j['heatmap']}
    assert 99 <= giugno[u_pri]['percentuale'] <= 101, giugno[u_pri]   # prima ~24%: colonna rossa
    assert 99 <= giugno[u_inf]['percentuale'] <= 101, giugno[u_inf]


def test_d2_confronto_annuale_senza_dati_nell_anno_recente_non_e_un_calo(client, db_mod):
    db = db_mod
    anni = ('2039-2040', '2040-2041')
    try:
        for a in anni:
            db.set_calendario(a, 3, int(a[5:]), 20, 20)
        _, uid = _commessa_con_utente(db, 'CONFRONTO VUOTO', nome='Confronto')
        _set_ore(db, uid, 2040, 3, 40)                        # marzo 2040 si', marzo 2041 no
        righe = client.get('/api/stats/confronto-annuale-dettaglio?mese=3').get_json()['confronto']
        recente = [r for r in righe if r['anno_scolastico'] == '2040-2041'][0]
        assert recente['num_utenti'] == 0
        assert recente['variazione'] is None, recente             # prima: -100% in rosso
        assert recente['nota']
    finally:
        with db.get_db_context() as conn:
            conn.execute('DELETE FROM calendario_scolastico WHERE anno_scolastico IN (?, ?)', anni)


def test_d3_colori_delle_commesse_distinti_per_grafici_e_dashboard():
    app_js = _leggi('static/js/app.js')
    assert 'function coloriCommesseDistinti(voci)' in app_js
    assert 'window.coloriCommesseDistinti = coloriCommesseDistinti' in app_js
    stat = _leggi('templates/statistiche.html')
    assert 'coloriCommesseDistinti(voci)' in stat and 'c.colore || null' not in stat
    dash = _leggi('static/js/dashboard.js')
    assert 'coloriCommesseDistinti(' in dash and "c.colore || 'var(--primary)'" not in dash


def test_d4_stato_del_mese_conta_solo_chi_e_in_servizio(client, db_mod):
    """Gia' corretto nella v1.19.0: qui si verifica che resti cosi'."""
    db = db_mod
    sid, dentro = _commessa_con_utente(db, 'STATO SERVIZIO', nome='Inservizio')
    fuori = db.get_or_create_utente(sid, 'Uscito', 'Prova', 10)
    db.update_utente_periodo(fuori, None, '2033-12')
    j = client.get(f'/api/stats/filtered?anno=2034&mese=1&commessa={_q("STATO SERVIZIO")}').get_json()
    assert j['num_utenti'] == 1
    da_completare = client.get('/api/stats/utenti-da-completare/2034/1').get_json()
    ids = {u['id'] if 'id' in u else u.get('utente_id') for u in da_completare}
    assert dentro in ids and fuori not in ids


# ==================== E1: Ctrl+Z ====================

def test_e1_azioni_vecchie_non_annullabili_e_ripulibili(client, db_mod):
    from datetime import datetime, timedelta
    db = db_mod
    _, uid = _commessa_con_utente(db, 'UNDO VECCHIO', nome='Vecchio')
    vecchio = (datetime.now() - timedelta(days=120)).isoformat()
    with db.get_db_context() as conn:
        conn.execute('INSERT INTO undo_actions (timestamp, action_type, data) VALUES (?, ?, ?)',
                     (vecchio, 'update_utente', json.dumps({'id': uid, 'dati_precedenti': {'monte_ore_settimanale': 3}})))
    try:
        j = client.get('/api/undo/ultima').get_json()
        assert j['azione'] is None and j['scadute'] >= 1 and j['ore_validita'] == 24
        r = client.post('/api/undo')
        assert r.status_code == 400 and '24 ore' in r.get_json()['error']
        assert 'quelle di più di 24 ore fa non si possono più annullare' in r.get_json()['error']
        assert db.get_utente_by_id(uid)['monte_ore_settimanale'] == 10          # non toccato
        assert client.get('/api/undo').get_json() == []
    finally:
        pulite = client.delete('/api/undo/scadute').get_json()
    assert pulite['rimosse'] >= 1
    atteso = ('Tolta 1 azione non più annullabile' if pulite['rimosse'] == 1
              else f"Tolte {pulite['rimosse']} azioni non più annullabili")
    assert atteso in [x['dettagli'] for x in db.get_audit_log(limit=20, entita='undo')]
    assert client.get('/api/undo/ultima').get_json()['scadute'] == 0


def test_e1_conferma_con_descrizione_e_data_poi_annulla(client, db_mod):
    db = db_mod
    _, uid = _commessa_con_utente(db, 'UNDO RECENTE', nome='Recente')
    assert client.put(f'/api/utenti/{uid}', json={'monte_ore': 12}).status_code == 200
    j = client.get('/api/undo/ultima').get_json()
    azione = j['azione']
    assert 'Recente Prova' in azione['descrizione'] and 'monte ore 12 → 10' in azione['descrizione']
    assert re.match(r'\d{2}/\d{2}/\d{4} alle \d{2}:\d{2}$', azione['quando'])
    # id diverso da quello mostrato nella conferma: non si annulla niente
    assert client.post('/api/undo', json={'id': azione['id'] + 1000}).status_code == 409
    assert db.get_utente_by_id(uid)['monte_ore_settimanale'] == 12
    assert client.post('/api/undo', json={'id': azione['id']}).status_code == 200
    assert db.get_utente_by_id(uid)['monte_ore_settimanale'] == 10
    audit = [a for a in client.get('/api/audit?entita=utente&limit=50').get_json()
             if a['azione'] == 'undo' and a['entita_id'] == uid]
    assert audit and 'monte_ore_settimanale' in (audit[0]['dati_nuovi'] or '')


def test_e1_la_pagina_chiede_conferma_e_ricarica_la_vista():
    app_js = _leggi('static/js/app.js')
    corpo = app_js.split('async function undoLastAction()', 1)[1].split('\n}\n', 1)[0]
    assert "/api/undo/ultima" in corpo and 'showConfirmDialog(' in corpo and 'azione.quando' in corpo
    assert corpo.index('showConfirmDialog(') < corpo.index("method: 'POST'")
    assert 'ricaricaVistaDopoAnnulla()' in corpo
    assert 'window.ricaricaDopoAnnulla = () => cambiaPeriodo()' in _rend_pagina()
