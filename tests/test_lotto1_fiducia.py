"""Lotto 1 "fiducia": bug silenziosi, blocco del mese chiuso, validazione degli
input, pulizia degli asset. Ogni test riproduce il problema originale e
verifica la correzione."""
import io
import os
from datetime import date

from openpyxl import load_workbook

PROJECT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def _leggi(path):
    with open(os.path.join(PROJECT_DIR, path), encoding='utf-8') as f:
        return f.read()


def _set_ore(db, uid, anno, mese, ore):
    db.get_or_create_rendicontazione(uid, anno, mese)
    db.update_rendicontazione(uid, anno, mese, ore_lavorate=ore)


def _celle(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if v is not None:
                    yield str(v)


def _testo_docx(data):
    from docx import Document
    doc = Document(io.BytesIO(data))
    testo = '\n'.join(p.text for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            testo += '\n' + ' | '.join(c.text for c in row.cells)
    return testo


# ---------- 1. storico monte ore ----------

def test_storico_monte_ore_registra_le_modifiche(client, db_mod):
    db = db_mod
    db.create_commessa('L1 STORICO')
    sid = db.get_or_create_scuola('L1 STORICO', 'IC L1 - Primaria')
    uid = db.get_or_create_utente(sid, 'Storico', 'Uno', 10)

    r = client.put(f'/api/utenti/{uid}', json={'monte_ore': 15})
    assert r.status_code == 200, r.data

    data = client.get(f'/api/utenti/{uid}/storico-monte-ore').get_json()
    assert data['utente']['monte_ore_attuale'] == 15
    assert data['storico'], 'storico sempre vuoto (dati salvati come repr Python, riletti come JSON)'
    ultimo = data['storico'][0]
    assert float(ultimo['monte_ore_precedente']) == 10
    assert float(ultimo['monte_ore_nuovo']) == 15


def test_storico_monte_ore_legge_anche_le_righe_storiche_repr(client, db_mod):
    """Le righe salvate prima della correzione (repr Python) restano leggibili."""
    db = db_mod
    db.create_commessa('L1 STORICO2')
    sid = db.get_or_create_scuola('L1 STORICO2', 'IC L1b - Primaria')
    uid = db.get_or_create_utente(sid, 'Storico', 'Due', 8)
    with db.get_db_context() as conn:
        conn.execute('''INSERT INTO audit_log (timestamp, azione, entita, entita_id, dettagli,
                        dati_precedenti, dati_nuovi) VALUES (?, ?, ?, ?, ?, ?, ?)''',
                     ('2025-01-01T10:00:00', 'modifica', 'utente', uid, 'Aggiornamento utente',
                      str({'monte_ore_settimanale': 8.0}), str({'monte_ore': 12})))
    data = client.get(f'/api/utenti/{uid}/storico-monte-ore').get_json()
    assert any(float(s['monte_ore_nuovo']) == 12 for s in data['storico'])


# ---------- 2. privacy (nomi puntati) anche su municipale e word ----------

def test_privacy_nomi_puntati_vale_anche_per_municipale_e_word(client, db_mod):
    db = db_mod
    db.create_commessa('L1 PRIVACY')
    sid = db.get_or_create_scuola('L1 PRIVACY', 'IC Priv - Primaria')
    uid = db.get_or_create_utente(sid, 'Riservato', 'Cognomelungo', 10)
    # poche ore (< 50% del previsto): cosi' l'utente compare anche nella
    # sezione "bassa erogazione" del Word, dove viene scritto il nome
    _set_ore(db, uid, 2025, 11, 5)

    r = client.get('/api/export/municipale/2025/11?commessa=L1%20PRIVACY')
    assert r.status_code == 200
    assert any('Riservato Cognomelungo' in c for c in _celle(load_workbook(io.BytesIO(r.data))))

    r = client.get('/api/export/municipale/2025/11?commessa=L1%20PRIVACY&privacy=true')
    assert r.status_code == 200
    assert 'privacy' in r.headers.get('Content-Disposition', '')
    assert not any('Cognomelungo' in c for c in _celle(load_workbook(io.BytesIO(r.data)))), \
        'nome in chiaro nel Riepilogo Municipale nonostante la spunta privacy'

    r = client.get('/api/export/word/2025/11?commessa=L1%20PRIVACY')
    assert r.status_code == 200
    assert 'Riservato Cognomelungo' in _testo_docx(r.data)
    r = client.get('/api/export/word/2025/11?commessa=L1%20PRIVACY&privacy=true')
    assert r.status_code == 200
    assert 'Cognomelungo' not in _testo_docx(r.data), 'nome in chiaro nel Word nonostante la privacy'


# ---------- 3. "esporta filtrati" ----------

def test_export_excel_rispetta_filtro_scuola_e_ricerca(client, db_mod):
    db = db_mod
    db.create_commessa('L1 FILTRI')
    s1 = db.get_or_create_scuola('L1 FILTRI', 'IC Filtri - Plesso Alfa')
    s2 = db.get_or_create_scuola('L1 FILTRI', 'IC Filtri - Plesso Beta')
    u1 = db.get_or_create_utente(s1, 'Alfa', 'Utente', 10)
    u2 = db.get_or_create_utente(s2, 'Beta', 'Utente', 10)
    _set_ore(db, u1, 2025, 11, 10)
    _set_ore(db, u2, 2025, 11, 12)

    r = client.get('/api/export/excel/2025/11?commessa=L1%20FILTRI'
                   '&scuola=IC%20Filtri%20-%20Plesso%20Alfa')
    assert r.status_code == 200
    celle = list(_celle(load_workbook(io.BytesIO(r.data))))
    assert any('Alfa Utente' in c for c in celle)
    assert not any('Beta Utente' in c for c in celle), 'il file "filtrato" conteneva tutti gli utenti'

    r = client.get('/api/export/excel/2025/11?commessa=L1%20FILTRI&search=beta')
    celle = list(_celle(load_workbook(io.BytesIO(r.data))))
    assert any('Beta Utente' in c for c in celle)
    assert not any('Alfa Utente' in c for c in celle)


# ---------- 5. undo dell'eliminazione ripristina TUTTO ----------

def test_undo_eliminazione_ripristina_variazioni_e_assegnazioni(client, db_mod):
    db = db_mod
    db.create_commessa('L1 UNDO')
    sid = db.get_or_create_scuola('L1 UNDO', 'IC Undo - Primaria')
    uid = db.get_or_create_utente(sid, 'Undo', 'Utente', 10)
    db.add_variazione_monte_ore(uid, 14, '2025-11', 'aumento')
    did = db.create_dipendente({'nome': 'Op', 'cognome': 'Eratore'})
    db.create_assegnazione(uid, did, 10)
    assert len(db.get_variazioni_monte_ore(uid)) == 1
    assert len(db.get_assegnazioni_utente(uid)) == 1

    assert client.delete(f'/api/utenti/{uid}').status_code == 200
    assert db.get_variazioni_monte_ore(uid) == []
    r = client.post('/api/undo')
    assert r.status_code == 200, r.data

    assert db.get_utente_by_id(uid) is not None
    variazioni = db.get_variazioni_monte_ore(uid)
    assert len(variazioni) == 1 and variazioni[0]['monte_ore'] == 14, "variazioni perse dopo l'undo"
    assert len(db.get_assegnazioni_utente(uid)) == 1, "assegnazioni perse dopo l'undo"


# ---------- 9. mese chiuso: blocca davvero ----------

def test_mese_chiuso_blocca_ogni_scrittura_delle_ore(client, db_mod):
    db = db_mod
    db.create_commessa('L1 CHIUSO')
    sid = db.get_or_create_scuola('L1 CHIUSO', 'IC Chiuso - Primaria')
    uid = db.get_or_create_utente(sid, 'Chiuso', 'Utente', 10)
    anno, mese = 2024, 2
    _set_ore(db, uid, anno, mese, 10)
    base = f'/api/rendicontazione/{anno}/{mese}'

    assert client.post(f'/api/mese-chiuso/{anno}/{mese}').status_code == 200
    try:
        r = client.post(base, json={'utente_id': uid, 'ore_lavorate_60': 20})
        assert r.status_code == 409 and r.get_json()['code'] == 'MESE_CHIUSO'
        r = client.post(base + '/batch', json={'updates': [{'utente_id': uid, 'ore_lavorate_60': 20}]})
        assert r.status_code == 409
        assert client.post(base + '/copia-precedente', json={}).status_code == 409
        assert client.post(base + '/compila-media', json={}).status_code == 409
        assert db.get_rendicontazione_completa(anno, mese, 'L1 CHIUSO')[0]['ore_lavorate_60'] == 10
        assert client.get(base + '?commessa=L1%20CHIUSO').get_json()['mese_chiuso']
    finally:
        client.delete(f'/api/mese-chiuso/{anno}/{mese}')

    # riaperto: la scrittura passa
    assert client.post(base, json={'utente_id': uid, 'ore_lavorate_60': 20}).status_code == 200
    assert db.get_rendicontazione_completa(anno, mese, 'L1 CHIUSO')[0]['ore_lavorate_60'] == 20


# ---------- 10. validazione degli input ----------

def test_validazione_ore_pasti_mese_e_utente(client, db_mod):
    db = db_mod
    db.create_commessa('L1 VALID')
    sid = db.get_or_create_scuola('L1 VALID', 'IC Valid - Primaria')
    uid = db.get_or_create_utente(sid, 'Valid', 'Utente', 10)
    base = '/api/rendicontazione/2025/11'

    assert client.post(base, json={'utente_id': uid, 'ore_lavorate_60': -5}).status_code == 400
    assert client.post(base, json={'utente_id': uid, 'ore_lavorate_60': 99999}).status_code == 400
    assert client.post(base, json={'utente_id': uid, 'pasti': 40}).status_code == 400
    assert client.post(base, json={'utente_id': uid, 'ore_lavorate_60': 'abc'}).status_code == 400
    assert client.post('/api/rendicontazione/2025/13',
                       json={'utente_id': uid, 'ore_lavorate_60': 1}).status_code == 400
    assert client.post(base, json={'utente_id': 999999, 'ore_lavorate_60': 1}).status_code == 404

    # batch: una riga sbagliata blocca tutto (nessuna scrittura parziale)
    r = client.post(base + '/batch', json={'updates': [
        {'utente_id': uid, 'ore_lavorate_60': 5},
        {'utente_id': uid, 'ore_lavorate_60': -1},
    ]})
    assert r.status_code == 400 and 'Riga 2' in r.get_json()['error']
    assert not any((d['ore_lavorate_60'] or 0) for d in db.get_rendicontazione_completa(2025, 11, 'L1 VALID'))

    # valori corretti passano
    assert client.post(base, json={'utente_id': uid, 'ore_lavorate_60': 12.5, 'pasti': 3}).status_code == 200
    riga = db.get_rendicontazione_completa(2025, 11, 'L1 VALID')[0]
    assert riga['ore_lavorate_60'] == 12.5 and riga['pasti'] == 3


def test_validazione_calendario_giorni_lavorativi(client):
    ok = {'anno_scolastico': '2031-2032', 'mese': 10, 'anno': 2031, 'giorni_lavorativi': 22}
    r = client.post('/api/calendario', json={**ok, 'giorni_lavorativi': 220})
    assert r.status_code == 400, 'giorni fuori scala accettati: media prevista x10 per tutti gli utenti'
    assert client.post('/api/calendario', json={'mese': 10}).status_code == 400
    assert client.post('/api/calendario', json={**ok, 'mese': 13}).status_code == 400
    assert client.post('/api/calendario', json={**ok, 'giorni_lavorativi_altri': 99}).status_code == 400
    assert client.post('/api/calendario', json={**ok, 'giorni_lavorativi_altri': ''}).status_code == 200


# ---------- 8. nessun anno scolastico scritto a mano ----------

def test_default_anno_scolastico_segue_la_data_di_oggi():
    import inspect
    import config
    import app as app_mod
    import routes_utenti_dettaglio as rud

    oggi = date.today()
    atteso = f'{oggi.year}-{oggi.year + 1}' if oggi.month >= 9 else f'{oggi.year - 1}-{oggi.year}'
    assert config.anno_scolastico_corrente() == atteso
    assert "'2025-2026'" not in inspect.getsource(app_mod)
    assert "'2025-2026'" not in inspect.getsource(rud)


# ---------- 21/22. niente font esterni, asset morti rimossi ----------

def test_nessuna_richiesta_a_font_esterni_e_asset_morti_rimossi():
    for t in ('templates/base.html', 'templates/login.html', 'templates/setup.html'):
        assert 'googleapis' not in _leggi(t), t
    assert '@import' not in _leggi('static/css/style.css')
    for f in ('static/js/advanced-features.js', 'static/js/search-advanced.js',
              'static/css/advanced-features.css', 'static/css/search-advanced.css',
              'static/images/report'):
        assert not os.path.exists(os.path.join(PROJECT_DIR, f)), f
    base = _leggi('templates/base.html')
    assert 'advanced-features' not in base and 'chart.umd.min.js' not in base
    # le regole ancora usate (heatmap, report rapidi, filtri avanzati) vivono in components.css
    css = _leggi('static/css/components.css')
    for cls in ('.heatmap-table', '.report-quick-card', '.report-quick-icon', '.filter-details', '.validation-item'):
        assert cls in css, cls
    assert 'components.css' in base
    for t in ('templates/index.html', 'templates/statistiche.html', 'templates/utente_dettaglio.html'):
        assert 'chart.umd.min.js' in _leggi(t), t
