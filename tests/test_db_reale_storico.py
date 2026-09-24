"""DB reale, parte 1: lo storico non deve cambiare per colpa del nuovo anno.

- A1 archiviare un utente non toglie le sue ore (ne' lui) dai mesi passati;
- A2 il passo Utenti del nuovo anno non propone archiviazioni gia' spuntate e
  l'esito resta visibile;
- A3 cambiare il monte ore nel passo Utenti non cambia i mesi passati, e le
  colonne "Monte Ore" degli export mostrano il monte ore effettivo del mese;
- A4 la finestra "Modifica utente" chiede conferma se si cambia il monte ore di
  chi ha gia' mesi rendicontati;
- F  le liste d'attesa valgono solo nel loro anno scolastico.

Solo dati inventati: nomi di fantasia e commesse di prova.
"""
import io
import json
import os
import re
from datetime import datetime

import docx
from openpyxl import load_workbook

import config

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MESI_2025_26 = [(2025, 9), (2025, 10), (2025, 11), (2025, 12), (2026, 1),
                (2026, 2), (2026, 3), (2026, 4), (2026, 5), (2026, 6)]


def _leggi(rel):
    with open(os.path.join(ROOT, rel), encoding='utf-8') as f:
        return f.read()


def _set_ore(db, uid, anno, mese, ore):
    db.get_or_create_rendicontazione(uid, anno, mese)
    db.update_rendicontazione(uid, anno, mese, ore_lavorate=ore)


def _mese_futuro():
    oggi = datetime.now()
    return (oggi.year + 1, oggi.month) if oggi.month < 12 else (oggi.year + 1, 1)


def _q(commessa):
    return commessa.replace(' ', '%20')


def _numeri_mese(client, anno, mese, commessa):
    """Ore, importi, utenti e ore previste di un mese dalla Rendicontazione."""
    j = client.get(f'/api/rendicontazione/{anno}/{mese}?commessa={_q(commessa)}').get_json()
    tg = j['totale_generale']
    return {'ore': round(tg['ore_lavorate_60'], 2), 'totale': tg['totale_100'], 'n': tg['num_utenti'],
            'cd': tg['credito_debito'], 'pasti': tg['pasti'],
            'prev': round(sum(d['media_con_assenza_60'] or 0 for d in j['dati']), 2),
            'ids': sorted(d['utente_id'] for d in j['dati'])}


def _celle_xlsx(data):
    wb = load_workbook(io.BytesIO(data), data_only=True)
    return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}


def _senza_data_generazione(celle):
    rumore = re.compile(r'generat', re.I)
    return {k: [[None if isinstance(c, str) and rumore.search(c) else c for c in r] for r in v]
            for k, v in celle.items()}


def _prepara_archivio(db, commessa):
    """Commessa con tre utenti: 'Resta' sempre attivo, 'Uscita' (data di fine a
    dicembre 2025, ore a ottobre e novembre, nessuna riga a dicembre) e 'Senzafine'
    (nessuna data di fine, ore solo a ottobre)."""
    db.create_commessa(commessa)
    sid = db.get_or_create_scuola(commessa, 'IC Fantasia - Primaria Arcobaleno')
    resta = db.get_or_create_utente(sid, 'Resta', 'Prova', 10)
    uscita = db.get_or_create_utente(sid, 'Uscita', 'Prova', 12)
    senzafine = db.get_or_create_utente(sid, 'Senzafine', 'Prova', 8)
    db.update_utente_periodo(uscita, None, '2025-12')
    for anno, mese in MESI_2025_26:
        _set_ore(db, resta, anno, mese, 30)
    _set_ore(db, uscita, 2025, 10, 44.5)
    _set_ore(db, uscita, 2025, 11, 0)       # riga a zero ore: conta comunque
    _set_ore(db, senzafine, 2025, 10, 21)
    return sid, resta, uscita, senzafine


# ==================== A1: archiviare non tocca il passato ====================

def test_archiviare_non_cambia_mesi_passati_ne_report(client, db_mod):
    db = db_mod
    commessa = 'STORICO ARCH'
    _, resta, uscita, senzafine = _prepara_archivio(db, commessa)
    stat_urls = [
        f'/api/stats/filtered?anno=2025&mese=10&commessa={_q(commessa)}',
        f'/api/stats/trend?anno_scolastico=2025-2026&commessa={_q(commessa)}',
        f'/api/stats/confronto-mese?anno=2025&mese=11&commessa={_q(commessa)}',
        f'/api/stats/top-scuole?anno=2025&mese=10&commessa={_q(commessa)}',
        f'/api/stats/scuole-dettaglio?anno=2025&mese=10&commessa={_q(commessa)}',
        f'/api/stats/validazione?anno=2025&mese=12&commessa={_q(commessa)}',
        f'/api/stats/heatmap/2025-2026?commessa={_q(commessa)}',
        f'/api/stats/ore-confronto/2025-2026?commessa={_q(commessa)}',
    ]
    export_urls = [f'/api/export/{t}/2025/{m}?commessa={_q(commessa)}'
                   for t in ('excel', 'municipale', 'dipartimentale') for m in (10, 11, 12)]
    export_urls.append(f'/api/export/annuale/2025-2026?commessa={_q(commessa)}')

    def fotografia():
        foto = {'mesi': {f'{a}-{m:02d}': _numeri_mese(client, a, m, commessa) for a, m in MESI_2025_26}}
        # con anno e mese num_utenti e' quello della vista mensile ("Stato del mese")
        foto['stats'] = {u: client.get(u).get_json() for u in stat_urls}
        # senza periodo e' il conteggio "di oggi": deve scendere
        foto['utenti_oggi'] = client.get(f'/api/stats/filtered?commessa={_q(commessa)}').get_json()['num_utenti']
        foto['export'] = {u: _senza_data_generazione(_celle_xlsx(client.get(u).data)) for u in export_urls}
        foto['da_completare'] = client.get('/api/stats/utenti-da-completare/2025/12').get_json()
        return foto

    prima = fotografia()
    assert uscita in prima['mesi']['2025-12']['ids']        # nel suo periodo, anche senza riga

    # archiviato con data di fine: ogni mese passato identico (ore, importi, utenti,
    # ore previste, credito/debito), statistiche ed export compresi
    assert client.put(f'/api/utenti/{uscita}', json={'attivo': False}).status_code == 200
    dopo = fotografia()
    assert dopo['mesi'] == prima['mesi']
    assert uscita in dopo['mesi']['2025-12']['ids']        # anche nel mese senza riga
    assert dopo['stats'] == prima['stats']
    assert dopo['export'] == prima['export']
    assert dopo['utenti_oggi'] == prima['utenti_oggi'] - 1
    assert any(u['id'] == uscita for u in dopo['da_completare'])

    # archiviato senza data di fine: nei mesi prima dell'archiviazione conta come
    # quando era attivo, anche dove non ha righe (utenti.archiviato_dal)
    assert client.put(f'/api/utenti/{senzafine}', json={'attivo': False}).status_code == 200
    dopo2 = fotografia()
    assert dopo2['mesi'] == prima['mesi'] and dopo2['stats'] == prima['stats']
    assert dopo2['export'] == prima['export']
    assert senzafine in dopo2['mesi']['2025-10']['ids'] and senzafine in dopo2['mesi']['2026-02']['ids']
    assert dopo2['utenti_oggi'] == prima['utenti_oggi'] - 2
    # mesi futuri senza righe: nessuno dei due
    fa, fm = _mese_futuro()
    ids_futuro = [d['utente_id'] for d in db.get_rendicontazione_completa(fa, fm, commessa)]
    assert resta in ids_futuro and uscita not in ids_futuro and senzafine not in ids_futuro

    # elenchi "di oggi": fuori dagli attivi, dentro gli archiviati
    attivi = [u['id'] for u in client.get(f'/api/utenti?commessa={_q(commessa)}').get_json()]
    archiviati = [u['id'] for u in client.get(f'/api/utenti?commessa={_q(commessa)}&attivi=0').get_json()]
    assert attivi == [resta] and set(archiviati) == {uscita, senzafine}

    # la Rendicontazione dice chi e' archiviato (per il segno "archiviato")
    righe = client.get(f'/api/rendicontazione/2025/10?commessa={_q(commessa)}').get_json()['dati']
    stato = {d['utente_id']: d['attivo'] for d in righe}
    assert stato[resta] == 1 and stato[uscita] == 0 and stato[senzafine] == 0


def test_archiviare_chi_ha_mesi_senza_righe_non_cambia_il_passato(client, db_mod):
    """Utente SENZA data di fine con mesi dell'anno senza riga (prima della prima
    riga e buchi in mezzo): archiviarlo non lo toglie da quei mesi. Prima contava
    solo dove aveva una riga e cambiavano utenti, ore previste e credito/debito."""
    db = db_mod
    commessa = 'STORICO BUCHI'
    db.create_commessa(commessa)
    sid = db.get_or_create_scuola(commessa, 'IC Fantasia - Primaria Buchi')
    resta = db.get_or_create_utente(sid, 'Resta', 'Buchi', 10)
    buchi = db.get_or_create_utente(sid, 'Buchi', 'Prova', 12)
    for anno, mese in MESI_2025_26:
        _set_ore(db, resta, anno, mese, 30)
    for anno, mese in [(2025, 10), (2026, 1), (2026, 5)]:       # niente set, nov, dic, feb...
        _set_ore(db, buchi, anno, mese, 40)
    export_urls = [f'/api/export/{t}/2025/{m}?commessa={_q(commessa)}'
                   for t in ('excel', 'municipale', 'dipartimentale') for m in (9, 11, 12)]
    export_urls.append(f'/api/export/annuale/2025-2026?commessa={_q(commessa)}')
    stat_urls = [f'/api/stats/filtered?anno=2025&mese=11&commessa={_q(commessa)}',
                 f'/api/stats/validazione?anno=2025&mese=11&commessa={_q(commessa)}',
                 f'/api/stats/heatmap/2025-2026?commessa={_q(commessa)}',
                 '/api/utenti/classifica/2025?mese=11&limit=500']

    def fotografia():
        return {'mesi': {f'{a}-{m:02d}': _numeri_mese(client, a, m, commessa) for a, m in MESI_2025_26},
                'stats': {u: client.get(u).get_json() for u in stat_urls},
                'export': {u: _senza_data_generazione(_celle_xlsx(client.get(u).data)) for u in export_urls},
                'word': [p.text for p in docx.Document(io.BytesIO(client.get(
                    f'/api/export/word/2025/11?commessa={_q(commessa)}').data)).paragraphs
                         if not p.text.startswith('Emesso il')]}

    prima = fotografia()
    assert buchi in prima['mesi']['2025-11']['ids'] and buchi in prima['mesi']['2025-09']['ids']
    assert client.put(f'/api/utenti/{buchi}', json={'attivo': False}).status_code == 200
    assert fotografia() == prima                  # tutto il 2025-26 identico
    assert db.get_utente_by_id(buchi)['archiviato_dal'] == db.mese_oggi()
    fa, fm = _mese_futuro()                       # mese futuro senza righe: non c'e'
    assert buchi not in [d['utente_id'] for d in db.get_rendicontazione_completa(fa, fm, commessa)]

    # salvare di nuovo l'archiviato non sposta il mese dell'archiviazione
    with db.get_db_context() as conn:
        conn.execute("UPDATE utenti SET archiviato_dal = '2026-02' WHERE id = ?", (buchi,))
    assert client.put(f'/api/utenti/{buchi}', json={'attivo': False, 'nome': 'Buchi'}).status_code == 200
    assert db.get_utente_by_id(buchi)['archiviato_dal'] == '2026-02'
    ids = {k: v['ids'] for k, v in
           {f'{a}-{m:02d}': _numeri_mese(client, a, m, commessa) for a, m in MESI_2025_26}.items()}
    assert buchi in ids['2026-01'] and buchi not in ids['2026-02'] and buchi in ids['2026-05']  # riga a maggio

    # ripristino: torna attivo e il mese dell'archiviazione sparisce; Ctrl+Z lo rimette
    assert client.put(f'/api/utenti/{buchi}', json={'attivo': True}).status_code == 200
    assert db.get_utente_by_id(buchi)['archiviato_dal'] is None
    client.post('/api/undo')
    u = db.get_utente_by_id(buchi)
    assert u['attivo'] == 0 and u['archiviato_dal'] == '2026-02'

    # archiviato da una versione precedente (senza mese): conta fino all'ultimo mese
    # in cui ha una riga, buchi compresi, e non dopo
    with db.get_db_context() as conn:
        conn.execute("UPDATE utenti SET archiviato_dal = NULL WHERE id = ?", (buchi,))
    ids = {k: v['ids'] for k, v in
           {f'{a}-{m:02d}': _numeri_mese(client, a, m, commessa) for a, m in MESI_2025_26}.items()}
    assert all(buchi in ids[k] for k in ('2025-09', '2025-11', '2026-02', '2026-04', '2026-05'))
    assert buchi not in ids['2026-06']

    # la migrazione aggiunge la colonna in modo idempotente
    db.init_db()
    db.init_db()
    assert 'archiviato_dal' in db.get_utente_by_id(buchi)


def test_archiviare_dal_passo_utenti_vale_dal_nuovo_anno(client, db_mod):
    """Il passo Utenti archivia dal primo mese del nuovo anno (o da oggi se piu'
    avanti): chi viene archiviato prima di settembre resta nei mesi in corso."""
    db = db_mod
    commessa = 'STORICO WIZ ARCH'
    db.create_commessa(commessa)
    sid = db.get_or_create_scuola(commessa, 'IC Fantasia - Primaria Futura')
    uid = db.get_or_create_utente(sid, 'Futuro', 'Prova', 10)
    r = client.post('/api/anno-scolastico/prepara-utenti', json={
        'anno_scolastico': '2040-2041', 'chiudi_variazioni': False, 'monte_ore': {}, 'archivia': [uid]})
    assert r.status_code == 200, r.data
    oggi = datetime.now()
    assert uid in [d['utente_id'] for d in db.get_rendicontazione_completa(oggi.year, oggi.month, commessa)]
    assert uid not in [d['utente_id'] for d in db.get_rendicontazione_completa(2040, 10, commessa)]
    u = db.get_utente_by_id(uid)
    assert u['attivo'] == 0 and u['archiviato_dal'] == '2040-09'
    assert uid not in [u['id'] for u in client.get(f'/api/utenti?commessa={_q(commessa)}').get_json()]
    # nell'anno corrente oggi non e' prima di settembre 2040: vale il mese di oggi
    assert db.mese_archiviazione(config.anno_scolastico_corrente()) == db.mese_oggi()


def test_stato_del_mese_conta_gli_utenti_della_vista_mensile(client, db_mod):
    """/api/stats/filtered con anno e mese: num_utenti e' il numero di utenti della
    vista mensile (chi era in servizio, archiviati compresi), cosi' "N di M utenti
    rendicontati" di un mese passato non cambia archiviando e non conta chi e' fuori
    periodo. Senza periodo resta il conteggio di oggi."""
    db = db_mod
    commessa = 'STORICO STATO MESE'
    db.create_commessa(commessa)
    sid = db.get_or_create_scuola(commessa, 'IC Fantasia - Primaria Stato')
    a = db.get_or_create_utente(sid, 'Alfa', 'Stato', 10)
    b = db.get_or_create_utente(sid, 'Beta', 'Stato', 10)
    uscito = db.get_or_create_utente(sid, 'Uscito', 'Stato', 10)
    db.update_utente_periodo(uscito, None, '2025-10')
    _set_ore(db, a, 2025, 11, 30)

    def numeri(anno, mese):
        tot = client.get(f'/api/stats/filtered?anno={anno}&mese={mese}&commessa={_q(commessa)}').get_json()
        da_fare = [u for u in client.get(f'/api/stats/utenti-da-completare/{anno}/{mese}').get_json()
                   if u['commessa'] == commessa]
        return tot['num_utenti'], len(da_fare)

    assert numeri(2025, 11) == (2, 1)            # 'Uscito' non era piu' in servizio
    assert client.put(f'/api/utenti/{b}', json={'attivo': False}).status_code == 200
    assert numeri(2025, 11) == (2, 1)            # mese passato invariato
    oggi = client.get(f'/api/stats/filtered?commessa={_q(commessa)}').get_json()['num_utenti']
    assert oggi == 2                              # utenti di oggi: Alfa e Uscito


def test_classifica_del_mese_segue_la_regola_degli_archiviati(client, db_mod):
    db = db_mod
    commessa = 'STORICO CLASSIFICA'
    db.create_commessa(commessa)
    sid = db.get_or_create_scuola(commessa, 'IC Fantasia - Primaria Classifica')
    uid = db.get_or_create_utente(sid, 'Classifica', 'Prova', 10)
    db.update_utente_periodo(uid, None, '2026-02')
    _set_ore(db, uid, 2025, 12, 20)              # nessuna riga a gennaio
    client.put(f'/api/utenti/{uid}', json={'attivo': False})
    ids = [r['id'] for r in client.get('/api/utenti/classifica/2026?mese=1&limit=5000').get_json()['classifica']]
    assert uid in ids                             # gennaio e' nel suo periodo di servizio
    # anno intero senza righe: c'e' (archiviato dopo gennaio, allora era attivo)
    ids = [r['id'] for r in client.get('/api/utenti/classifica/2026?limit=5000').get_json()['classifica']]
    assert uid in ids


def test_heatmap_di_un_anno_passato_conta_anche_gli_archiviati(client, db_mod):
    """Statistiche > heatmap di un anno passato: il totale ("Mostrati N su ...") e il
    limite di "Mostra tutti" sono gli utenti di QUELL'anno, archiviati compresi, non
    gli attivi di oggi. Prima la pagina usava gli attivi di oggi: archiviando un
    utente la heatmap 2025-26 tagliava in silenzio l'ultimo utente NON archiviato."""
    db = db_mod
    commessa = 'STORICO HEATMAP'
    db.create_commessa(commessa)
    sid = db.get_or_create_scuola(commessa, 'IC Fantasia - Primaria Mappa')
    ids = [db.get_or_create_utente(sid, nome, 'Mappa', 10) for nome in ('Alfa', 'Beta', 'Gamma', 'Zeta')]
    for uid in ids:
        _set_ore(db, uid, 2025, 10, 30)
    url = f'/api/stats/heatmap/2025-2026?commessa={_q(commessa)}'
    prima = client.get(url).get_json()
    assert prima['totale'] == 4 and len(prima['heatmap']) == 4

    assert client.put(f'/api/utenti/{ids[0]}', json={'attivo': False}).status_code == 200
    oggi = client.get(f'/api/stats/advanced?commessa={_q(commessa)}').get_json()['num_utenti']
    assert oggi == 3                                   # attivi di oggi
    dopo = client.get(url).get_json()
    assert dopo['totale'] == 4                         # l'anno 2025-26 li ha tutti e quattro
    assert dopo['heatmap'] == prima['heatmap']
    # con il limite della pagina ("Mostra tutti" = il totale) c'e' anche l'ultimo, Zeta
    tutti = client.get(f'{url}&limit={dopo["totale"]}').get_json()
    assert sorted(u['id'] for u in tutti['heatmap']) == sorted(ids)
    # limite piu' basso: il totale resta quello dell'anno
    parziale = client.get(f'{url}&limit=2').get_json()
    assert len(parziale['heatmap']) == 2 and parziale['totale'] == 4

    # la pagina usa il totale dell'anno restituito dall'API, non quello di oggi
    pagina = _leggi('templates/statistiche.html')
    assert 'totaleUtenti = data.totale' in pagina


def test_copia_e_compila_non_scrivono_righe_per_gli_archiviati(client, db_mod):
    db = db_mod
    commessa = 'STORICO COPIA'
    _, resta, uscita, senzafine = _prepara_archivio(db, commessa)
    for uid in (uscita, senzafine):
        client.put(f'/api/utenti/{uid}', json={'attivo': False})

    # novembre <- ottobre: 'Senzafine' e 'Uscita' hanno ore a ottobre ma sono archiviati
    # (solo gli utenti di questa prova: la copia senza elenco vale per tutti)
    r = client.post('/api/rendicontazione/2025/11/copia-precedente',
                    json={'solo_vuoti': False, 'utente_ids': [resta, uscita, senzafine]})
    assert r.status_code == 200
    with db.get_db_context() as conn:
        righe = {row[0] for row in conn.execute(
            'SELECT utente_id FROM rendicontazione WHERE anno = 2025 AND mese = 11 AND utente_id IN (?, ?)',
            (senzafine, uscita)).fetchall()}
    assert senzafine not in righe            # nessuna riga nuova per l'archiviato
    # dicembre: 'Uscita' e' nel suo periodo (compare) ma non si compila
    assert uscita in [d['utente_id'] for d in db.get_rendicontazione_completa(2025, 12, commessa)]
    r = client.post('/api/rendicontazione/2025/12/compila-media', json={'utente_ids': [uscita]})
    assert r.status_code == 200
    with db.get_db_context() as conn:
        assert conn.execute('SELECT COUNT(*) FROM rendicontazione WHERE utente_id = ? AND anno = 2025 AND mese = 12',
                            (uscita,)).fetchone()[0] == 0


def test_testi_di_archiviazione_dicono_il_vero():
    utenti = _leggi('templates/utenti.html')
    dashboard = _leggi('static/js/dashboard.js')
    assert 'tutto lo storico resta' not in utenti
    assert 'Non comparirà nei mesi futuri' in utenti and 'i mesi già rendicontati restano invariati' in utenti
    assert 'non compariranno nei mesi futuri; i mesi già rendicontati restano' in dashboard


# ==================== A2: passo Utenti senza archiviazioni proposte ====================

def test_wizard_non_propone_archiviazioni_gia_spuntate(client, db_mod):
    db = db_mod
    commessa = 'STORICO WIZ A2'
    _, resta, uscita, _ = _prepara_archivio(db, commessa)
    r = client.get('/api/anno-scolastico/utenti-anteprima?anno_scolastico=2026-2027')
    assert r.status_code == 200
    body = r.get_json()
    ant = {u['id']: u for u in body['utenti']}
    assert ant[uscita]['uscito'] is True and ant[uscita]['proposta_archivio'] is False
    assert ant[resta]['uscito'] is False
    assert body['da_archiviare'] == 0 and body['usciti'] >= 1

    js = _leggi('static/js/dashboard.js')
    assert 'u.proposta_archivio' not in js           # nessuna spunta proposta dal server
    assert "'uscito a'" in js or 'uscito a' in js    # solo l'indicazione "uscito a mm/aaaa"


def test_esito_del_nuovo_anno_resta_visibile():
    """loadBannerNuovoAnno non nasconde il riquadro se c'e' un esito appena scritto
    (prima spariva subito, sia dopo il passo 1 sia dopo il passo 2)."""
    js = _leggi('static/js/dashboard.js')
    corpo = js[js.index('async function loadBannerNuovoAnno'):js.index('function _statoWizard')]
    blocco = corpo[corpo.index('if (!stato.mostra_banner)'):corpo.index("card.style.display = 'none'")]
    assert "nuovo-anno-esito" in blocco and 'renderPassiNuovoAnno(stato)' in blocco


# ==================== A3: monte ore del nuovo anno e mesi passati ====================

def test_wizard_monte_ore_non_cambia_i_mesi_passati(client, db_mod):
    db = db_mod
    commessa = 'STORICO MONTE'
    db.create_commessa(commessa)
    sid = db.get_or_create_scuola(commessa, 'IC Fantasia - Primaria Girasole')
    semplice = db.get_or_create_utente(sid, 'Semplice', 'Prova', 6)       # nessuna variazione
    variato = db.get_or_create_utente(sid, 'Variato', 'Prova', 10)        # variazione dal primo mese
    tardi = db.get_or_create_utente(sid, 'Tardi', 'Prova', 9)             # prima riga a novembre
    for anno, mese in MESI_2025_26:
        _set_ore(db, semplice, anno, mese, 25)
        _set_ore(db, variato, anno, mese, 40)
    db.add_variazione_monte_ore(variato, 14, '2025-09', 'aumento da inizio anno', mese_fine='2026-01')
    for anno, mese in MESI_2025_26[2:]:
        _set_ore(db, tardi, anno, mese, 30)

    anno_url = f'/api/export/annuale/2025-2026?commessa={_q(commessa)}'
    prima_mesi = {f'{a}-{m:02d}': _numeri_mese(client, a, m, commessa) for a, m in MESI_2025_26}
    prima_eff = {(a, m): {d['utente_id']: d['monte_ore_effettivo'] for d in db.get_rendicontazione_completa(a, m, commessa)}
                 for a, m in MESI_2025_26}
    prima_annuale = _senza_data_generazione(_celle_xlsx(client.get(anno_url).data))
    prima_excel = _senza_data_generazione(_celle_xlsx(client.get(f'/api/export/excel/2025/10?commessa={_q(commessa)}').data))

    r = client.post('/api/anno-scolastico/prepara-utenti', json={
        'anno_scolastico': '2026-2027', 'chiudi_variazioni': True,
        'monte_ore': {str(semplice): 8, str(variato): 12, str(tardi): 11}, 'archivia': []})
    assert r.status_code == 200, r.data
    assert r.get_json()['variazioni_conservate'] == 3

    # mesi passati identici: ore previste, credito/debito, monte ore effettivo, export
    assert {f'{a}-{m:02d}': _numeri_mese(client, a, m, commessa) for a, m in MESI_2025_26} == prima_mesi
    assert {(a, m): {d['utente_id']: d['monte_ore_effettivo'] for d in db.get_rendicontazione_completa(a, m, commessa)}
            for a, m in MESI_2025_26} == prima_eff
    assert _senza_data_generazione(_celle_xlsx(client.get(anno_url).data)) == prima_annuale
    assert _senza_data_generazione(_celle_xlsx(
        client.get(f'/api/export/excel/2025/10?commessa={_q(commessa)}').data)) == prima_excel
    # 'Tardi' compare (senza data di inizio) anche a settembre, prima della sua prima riga
    assert prima_eff[(2025, 9)][tardi] == 9
    # da settembre il valore nuovo
    nuovo = {d['utente_id']: d['monte_ore_effettivo'] for d in db.get_rendicontazione_completa(2026, 9, commessa)}
    assert nuovo[semplice] == 8 and nuovo[variato] == 12 and nuovo[tardi] == 11

    # la variazione registrata e' chiusa ad agosto e inizia PRIMA di quella esistente
    vs = db.get_variazioni_monte_ore(variato)
    conservata = [v for v in vs if v['monte_ore'] == 10][0]
    assert conservata['mese_fine'] == '2026-08' and conservata['mese_inizio'] < '2025-09'
    assert '31/08/2026' in (conservata['nota'] or '')

    # secondo anno di fila (8 -> 10): ogni anno tiene il suo valore
    r = client.post('/api/anno-scolastico/prepara-utenti', json={
        'anno_scolastico': '2027-2028', 'monte_ore': {str(semplice): 10}, 'archivia': []})
    assert r.status_code == 200
    vs = db.get_variazioni_monte_ore(semplice)
    assert db.risolvi_monte_ore(10, vs, '2025-10') == 6
    assert db.risolvi_monte_ore(10, vs, '2026-10') == 8
    assert db.risolvi_monte_ore(10, vs, '2027-10') == 10


def test_wizard_diminuzione_con_data_inizio_non_crea_incrementi(client, db_mod):
    """Diminuire nel passo Utenti il monte ore di chi ha una data di inizio a meta'
    anno: la variazione che conserva il valore vecchio parte dal primo mese del
    gestionale, non da data_inizio. Prima a settembre (riferimento della colonna
    "Di cui hanno ricevuto incremento ore" del Municipale) valeva gia' il valore
    nuovo, piu' basso, e i mesi dopo data_inizio risultavano un "incremento"."""
    db = db_mod
    commessa = 'STORICO DIMINUZIONE'
    db.create_commessa(commessa)
    sid = db.get_or_create_scuola(commessa, 'IC Fantasia - Primaria Diminuzione')
    sempre = db.get_or_create_utente(sid, 'Sempre', 'Prova', 10)
    tardivo = db.get_or_create_utente(sid, 'Tardivo', 'Prova', 8)
    db.update_utente_periodo(tardivo, '2026-01', None)
    db.update_utente_lista_attesa(tardivo, 'Gennaio', '2025-2026')
    for anno, mese in MESI_2025_26:
        _set_ore(db, sempre, anno, mese, 30)
        if (anno, mese) >= (2026, 1):
            _set_ore(db, tardivo, anno, mese, 28)
    mesi = [(2026, 1), (2026, 3), (2026, 6)]

    def fotografia():
        foto = {f'municipale {a}-{m}': _senza_data_generazione(_celle_xlsx(client.get(
            f'/api/export/municipale/{a}/{m}?commessa={_q(commessa)}').data)) for a, m in mesi}
        for a, m in mesi:
            d = docx.Document(io.BytesIO(client.get(f'/api/export/word/{a}/{m}?commessa={_q(commessa)}').data))
            foto[f'word {a}-{m}'] = [[c.text for c in r.cells] for t in d.tables for r in t.rows]
        # heatmap: ore e ore previste di ogni mese (il campo monte_ore e' la base di oggi)
        foto['heatmap'] = {u['id']: u['mesi'] for u in client.get(
            f'/api/stats/heatmap/2025-2026?commessa={_q(commessa)}').get_json()['heatmap']}
        return foto

    prima = fotografia()
    r = client.post('/api/anno-scolastico/prepara-utenti', json={
        'anno_scolastico': '2026-2027', 'chiudi_variazioni': False,
        'monte_ore': {str(tardivo): 5, str(sempre): 7}, 'archivia': []})
    assert r.status_code == 200, r.data
    assert fotografia() == prima
    conservata = [v for v in db.get_variazioni_monte_ore(tardivo) if v['monte_ore'] == 8][0]
    assert conservata['mese_inizio'] <= '2025-09' and conservata['mese_fine'] == '2026-08'
    assert db.get_monte_ore_effettivo_bulk(2025, 9)[tardivo] == 8
    nuovo = {d['utente_id']: d['monte_ore_effettivo'] for d in db.get_rendicontazione_completa(2026, 10, commessa)}
    assert nuovo[tardivo] == 5 and nuovo[sempre] == 7


def test_validazione_usa_il_monte_ore_del_mese(client, db_mod):
    """Dopo il passo Utenti che porta il monte ore base a 0, i mesi passati (monte
    ore effettivo di allora > 0) non mostrano la falsa anomalia "monte ore = 0" e
    l'utente resta nei controlli "senza ore" e "differenze"."""
    db = db_mod
    commessa = 'STORICO VALIDAZIONE'
    db.create_commessa(commessa)
    sid = db.get_or_create_scuola(commessa, 'IC Fantasia - Primaria Validazione')
    uid = db.get_or_create_utente(sid, 'Azzerato', 'Prova', 12)
    _set_ore(db, uid, 2025, 10, 3)               # molte meno ore delle previste
    urls = [f'/api/stats/validazione?anno={a}&mese={m}&commessa={_q(commessa)}'
            for a, m in [(2025, 10), (2025, 11)]]

    def categorie():
        return {u: {a['categoria']: sorted(str(d.get('id', d.get('nome'))) for d in a['dettagli'])
                    for a in client.get(u).get_json()['anomalie']} for u in urls}

    prima = categorie()
    assert str(uid) in prima[urls[0]]['differenze_elevate'] and str(uid) in prima[urls[1]]['ore_mancanti']
    r = client.post('/api/anno-scolastico/prepara-utenti', json={
        'anno_scolastico': '2026-2027', 'chiudi_variazioni': False, 'monte_ore': {str(uid): 0}, 'archivia': []})
    assert r.status_code == 200, r.data
    dopo = categorie()
    assert dopo == prima
    assert 'monte_ore_zero' not in dopo[urls[0]]


def test_colonne_monte_ore_degli_export_usano_il_monte_ore_del_mese(client, db_mod):
    db = db_mod
    commessa = 'STORICO COLONNE'
    db.create_commessa(commessa)
    sid = db.get_or_create_scuola(commessa, 'IC Fantasia - Primaria Colonne')
    uid = db.get_or_create_utente(sid, 'Colonna', 'Prova', 10)
    db.add_variazione_monte_ore(uid, 16, '2026-03', 'aumento')
    _set_ore(db, uid, 2026, 3, 50)
    _set_ore(db, uid, 2026, 2, 40)

    def colonna_monte_ore(url, foglio, intestazione):
        celle = _celle_xlsx(client.get(url).data)[foglio]
        riga_h = next(i for i, r in enumerate(celle) if intestazione in r)
        col = celle[riga_h].index(intestazione)
        return [r[col] for r in celle[riga_h + 1:] if r[col] is not None and any('Colonna' in str(c) for c in r if c)]

    assert colonna_monte_ore(f'/api/export/excel/2026/3?commessa={_q(commessa)}', 'Dettaglio', 'Monte Ore') == [16]
    assert colonna_monte_ore(f'/api/export/excel/2026/2?commessa={_q(commessa)}', 'Dettaglio', 'Monte Ore') == [10]
    assert colonna_monte_ore(f'/api/export/municipale/2026/3?commessa={_q(commessa)}',
                             'Dettaglio Utenti', 'Monte Ore') == [16]


# ==================== A4: "Modifica utente" e mesi gia' rendicontati ====================

def test_modifica_utente_chiede_conferma_se_ci_sono_mesi_rendicontati():
    html = _leggi('templates/utenti.html')
    corpo = html[html.index('function saveEditUtente'):html.index('// === Elimina utente ===')]
    assert 'confermaMonteOreTuttiIMesi' in corpo
    assert "confirmText: 'Cambia per tutti i mesi'" in corpo and "cancelText: 'Annulla'" in corpo
    assert "extraText: 'Apri le Variazioni'" in corpo and 'openVariazioniMonteOre()' in corpo
    assert '/storico-ore' in corpo                      # solo se ci sono mesi rendicontati
    assert 'Di cui hanno ricevuto incremento ore' in corpo
    # il terzo pulsante della conferma esiste davvero in showConfirmDialog
    app_js = _leggi('static/js/app.js')
    assert 'confirm-dialog-extra' in app_js and 'onExtra' in app_js


# ==================== F: liste d'attesa legate all'anno scolastico ====================

def test_etichette_esistenti_prendono_l_anno_dell_ultimo_mese(db_mod):
    db = db_mod
    commessa = 'STORICO LISTE MIG'
    db.create_commessa(commessa)
    sid = db.get_or_create_scuola(commessa, 'IC Fantasia - Primaria Liste')
    con_ore = db.get_or_create_utente(sid, 'Lista', 'Conore', 10)
    senza_ore = db.get_or_create_utente(sid, 'Lista', 'Senzaore', 10)
    vuota = db.get_or_create_utente(sid, 'Lista', 'Vuota', 10)
    _set_ore(db, con_ore, 2024, 11, 10)
    _set_ore(db, con_ore, 2026, 3, 12)          # ultimo mese: marzo 2026 -> 2025-2026
    with db.get_db_context() as conn:           # come un DB della versione vecchia
        conn.execute("UPDATE utenti SET lista_attesa = 'Novembre', lista_attesa_as = NULL WHERE id IN (?, ?)",
                     (con_ore, senza_ore))
        conn.execute("UPDATE utenti SET lista_attesa = '  ', lista_attesa_as = '2025-2026' WHERE id = ?", (vuota,))
    db.init_db()                                 # migrazione idempotente all'avvio
    db.init_db()
    assert db.get_utente_by_id(con_ore)['lista_attesa_as'] == '2025-2026'
    assert db.get_utente_by_id(senza_ore)['lista_attesa_as'] == config.anno_scolastico_corrente()
    assert db.get_utente_by_id(vuota)['lista_attesa_as'] is None


def test_municipale_e_word_contano_la_lista_solo_nel_suo_anno(client, db_mod):
    db = db_mod
    commessa = 'STORICO LISTE REP'
    db.create_commessa(commessa)
    sid = db.get_or_create_scuola(commessa, 'IC Fantasia - Primaria Report')
    nov = db.get_or_create_utente(sid, 'Novembrina', 'Prova', 10)
    altro = db.get_or_create_utente(sid, 'Altro', 'Prova', 10)
    db.update_utente_lista_attesa(nov, 'Novembre', '2025-2026')
    for uid in (nov, altro):
        _set_ore(db, uid, 2025, 11, 20)
        _set_ore(db, uid, 2026, 10, 22)

    def intestazioni_liste(anno, mese):
        ws = load_workbook(io.BytesIO(client.get(
            f'/api/export/municipale/{anno}/{mese}?commessa={_q(commessa)}').data), data_only=True)['Riepilogo Municipale']
        righe = list(ws.iter_rows(values_only=True))
        i = next(i for i, r in enumerate(righe) if r and r[0] and 'RIEPILOGATIVO PER LISTA' in str(r[0]))
        return [h for h in righe[i + 1][4:] if h], righe[i + 2]

    liste, alunni = intestazioni_liste(2025, 11)
    assert liste == ['Lista Nov 2025'] and alunni[2] == 1        # 1 non in lista, 1 in lista
    liste, alunni = intestazioni_liste(2026, 10)
    assert liste == [] and alunni[2] == 2                         # nel 2026-27 non conta piu'

    def testo_word(anno, mese):
        d = docx.Document(io.BytesIO(client.get(f'/api/export/word/{anno}/{mese}?commessa={_q(commessa)}').data))
        return '\n'.join(p.text for p in d.paragraphs)
    assert 'Di questi, 1 utenti risultano in lista di attesa' in testo_word(2025, 11)
    assert 'risultano in lista di attesa' not in testo_word(2026, 10)

    # nel dettaglio dell'Excel mensile la colonna Lista Attesa segue la stessa regola
    celle = _celle_xlsx(client.get(f'/api/export/excel/2026/10?commessa={_q(commessa)}').data)['Dettaglio']
    assert not any('Novembre' in [str(c) for c in r] for r in celle)


def test_ogni_scrittura_della_lista_imposta_l_anno(client, db_mod):
    db = db_mod
    commessa = 'STORICO LISTE SCR'
    db.create_commessa(commessa)
    sid = db.get_or_create_scuola(commessa, 'IC Fantasia - Primaria Scritture')
    uid = db.get_or_create_utente(sid, 'Scrittura', 'Prova', 10)
    corrente = config.anno_scolastico_corrente()

    # scheda utente (PUT): anno corrente, oppure quello indicato; senza etichetta niente anno
    assert client.put(f'/api/utenti/{uid}', json={'lista_attesa': 'Marzo'}).status_code == 200
    assert db.get_utente_by_id(uid)['lista_attesa_as'] == corrente
    assert client.put(f'/api/utenti/{uid}', json={'lista_attesa': 'Aprile', 'lista_attesa_as': '2025-2026'}).status_code == 200
    assert db.get_utente_by_id(uid)['lista_attesa_as'] == '2025-2026'
    assert client.put(f'/api/utenti/{uid}', json={'lista_attesa_as': 'x', 'lista_attesa': 'Maggio'}).status_code == 400
    # una modifica che non tocca la lista lascia l'anno com'era
    assert client.put(f'/api/utenti/{uid}', json={'monte_ore': 11}).status_code == 200
    assert db.get_utente_by_id(uid)['lista_attesa_as'] == '2025-2026'
    # annulla (Ctrl+Z) dell'ultima modifica e poi di quella della lista: torna anche l'anno
    client.post('/api/undo')
    client.post('/api/undo')
    assert db.get_utente_by_id(uid)['lista_attesa'] == 'Marzo'
    assert db.get_utente_by_id(uid)['lista_attesa_as'] == corrente
    assert client.put(f'/api/utenti/{uid}', json={'lista_attesa': None}).status_code == 200
    assert db.get_utente_by_id(uid)['lista_attesa_as'] is None

    # modifica massiva e suo annullamento
    r = client.put('/api/utenti/bulk', json={'updates': [{'id': uid, 'lista_attesa': 'Gennaio'}]})
    assert r.status_code == 200
    assert db.get_utente_by_id(uid)['lista_attesa_as'] == corrente
    client.post('/api/undo')
    assert db.get_utente_by_id(uid)['lista_attesa'] is None and db.get_utente_by_id(uid)['lista_attesa_as'] is None

    # duplica: la copia tiene etichetta e anno
    db.update_utente_lista_attesa(uid, 'Febbraio', '2020-2021')
    nuovo = client.post(f'/api/utenti/{uid}/duplica').get_json()['nuovo_id']
    assert db.get_utente_by_id(nuovo)['lista_attesa_as'] == '2020-2021'

    # esportazione CSV: l'anno in una colonna in fondo
    csv_txt = client.get(f'/api/utenti/export-csv?commessa={_q(commessa)}').data.decode('utf-8')
    intestazione = csv_txt.splitlines()[0].split(',')
    assert intestazione[-1] == 'lista_attesa_anno' and intestazione[5] == 'lista_attesa'
    assert '2020-2021' in csv_txt

    # pagina Utenti e scheda: etichetta di un anno passato -> non conta quest'anno,
    # si mostra con l'anno breve; quella di quest'anno conta
    utenti = {u['id']: u for u in client.get(f'/api/utenti?commessa={_q(commessa)}').get_json()}
    assert utenti[uid]['lista_attesa_corrente'] is False and utenti[uid]['lista_attesa_anno_breve'] == '20/21'
    dettaglio = client.get(f'/api/utente/{uid}/dettaglio').get_json()['utente']
    assert dettaglio['lista_attesa_corrente'] is False and dettaglio['lista_attesa_anno_breve'] == '20/21'
    client.put(f'/api/utenti/{uid}', json={'lista_attesa': 'Febbraio'})
    utenti = {u['id']: u for u in client.get(f'/api/utenti?commessa={_q(commessa)}').get_json()}
    assert utenti[uid]['lista_attesa_corrente'] is True
    # la pagina mostra attenuata (con l'anno) l'etichetta di un anno passato
    html = _leggi('templates/utenti.html')
    assert 'lista-anno-passato' in html and 'u.lista_attesa_corrente' in html


def test_trasloco_json_vecchio_ricava_l_anno_dalle_ore(client, db_mod):
    db = db_mod
    commessa = 'STORICO LISTE JSON'
    # file 2.0 (versione vecchia): etichetta senza anno, ore di aprile 2026
    dati = {
        'versione': '2.0',
        'commesse': [{'id': 901, 'nome': commessa, 'descrizione': None, 'colore': '#3B82F6', 'attiva': 1}],
        'scuole': [{'id': 902, 'commessa_id': 901, 'commessa_nome': commessa,
                    'nome_completo': 'IC Fantasia - Primaria Trasloco'}],
        'utenti': [{'id': 903, 'scuola_id': 902, 'scuola_nome': 'IC Fantasia - Primaria Trasloco',
                    'commessa_nome': commessa, 'nome': 'Trasloco', 'cognome': 'Prova', 'nome_puntato': 'T. P.',
                    'monte_ore_settimanale': 10, 'lista_attesa': 'Aprile', 'attivo': 1,
                    'data_inserimento': '2026-04-01T10:00:00'}],
        'rendicontazione': [{'id': 904, 'utente_id': 903, 'utente_nome': 'Trasloco', 'utente_cognome': 'Prova',
                             'anno': 2026, 'mese': 4, 'ore_lavorate_60': 30, 'pasti': 0,
                             'giorni_lavorativi': 17, 'note': None}],
        'calendario': [],
    }
    r = client.post('/api/migrazione/importa', data={
        'mode': 'merge', 'file': (io.BytesIO(json.dumps(dati).encode()), 'vecchio.json')},
        content_type='multipart/form-data')
    assert r.status_code == 200, r.data
    uid = [u['id'] for u in client.get(f'/api/utenti?commessa={_q(commessa)}').get_json()][0]
    assert db.get_utente_by_id(uid)['lista_attesa_as'] == '2025-2026'


def test_trasloco_json_porta_il_mese_dell_archiviazione(client, db_mod):
    """Esporta/importa (trasloco JSON): il mese dell'archiviazione viaggia col file;
    un file di una versione vecchia (senza) non cancella quello gia' presente."""
    db = db_mod
    commessa = 'STORICO ARCH JSON'
    utente = {'id': 913, 'scuola_id': 912, 'scuola_nome': 'IC Fantasia - Primaria Archivio',
              'commessa_nome': commessa, 'nome': 'Archivio', 'cognome': 'Json', 'nome_puntato': 'A. J.',
              'monte_ore_settimanale': 10, 'lista_attesa': None, 'attivo': 0, 'archiviato_dal': '2026-02',
              'data_inserimento': '2025-09-01T10:00:00'}

    def importa(u):
        dati = {'versione': '2.0',
                'commesse': [{'id': 911, 'nome': commessa, 'descrizione': None, 'colore': '#3B82F6', 'attiva': 1}],
                'scuole': [{'id': 912, 'commessa_id': 911, 'commessa_nome': commessa,
                            'nome_completo': 'IC Fantasia - Primaria Archivio'}],
                'utenti': [u], 'rendicontazione': [], 'calendario': []}
        r = client.post('/api/migrazione/importa', data={
            'mode': 'merge', 'file': (io.BytesIO(json.dumps(dati).encode()), 'trasloco.json')},
            content_type='multipart/form-data')
        assert r.status_code == 200, r.data
        return [u['id'] for u in client.get(f'/api/utenti?commessa={_q(commessa)}&attivi=0').get_json()]

    uid = importa(utente)[0]
    assert db.get_utente_by_id(uid)['archiviato_dal'] == '2026-02'
    vecchio = {k: v for k, v in utente.items() if k != 'archiviato_dal'}
    importa(vecchio)
    assert db.get_utente_by_id(uid)['archiviato_dal'] == '2026-02'
    esportato = json.loads(client.get('/api/migrazione/esporta').data)
    assert [u for u in esportato['utenti'] if u['id'] == uid][0]['archiviato_dal'] == '2026-02'
    importa(dict(utente, attivo=1))
    assert db.get_utente_by_id(uid)['archiviato_dal'] is None and db.get_utente_by_id(uid)['attivo'] == 1
